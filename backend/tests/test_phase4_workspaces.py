"""Phase 4B: workspaces and board packs (plan §10, stream 4B).

Pinned here, without a database (fake sessions compile every statement for PostgreSQL):

1. **Landing** — line of defence from role names then permissions; the workspace each
   line starts on; a start-page choice wins while it is still available.
2. **Snapshots** — when one is due, which snapshot stands for a date, the rows a capture
   writes, reconstruction from history, the trend series.
3. **Assurance** — findings by age and owner, and the three-lines map with its gaps.
4. **Board packs** — draft → reviewed → released with segregation of duties, commentary
   rules, stored figures round-trip, saved section order, past-period snapshot figures,
   charts and branding in the PDF.
5. **My Work** — the new kinds build valid SQL and match free-text people.
6. **Roles** — the Board Member role arrives once for existing organisations.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta, timezone
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from sqlalchemy.dialects import postgresql

from app.services import board_pack as bp
from app.services import my_work as mw
from app.services import snapshots as sn
from app.services import workspaces as ws
from app.services.notifications import Directory, DirectoryUser

TODAY = date(2026, 9, 17)
KHI = ZoneInfo("Asia/Karachi")


# =================================================================== landing ===
@pytest.mark.parametrize("roles,perms,line", [
    (["Internal Auditor"], ["risk:read"], "audit"),
    (["Branch Operations"], ["risk:write"], "first_line"),
    (["Risk Manager"], ["risk:read"], "second_line"),
    (["Board Member"], ["board:read"], "board"),
    ([], ["risk:read", "board:read"], "board"),                      # nothing to write
    ([], ["internal_audit:write", "risk:read"], "audit"),
    ([], ["internal_audit:write", "risk:write"], "second_line"),
    ([], ["issue:write"], "first_line"),
])
def test_line_of_defence(roles, perms, line):
    assert ws.line_of_defence(roles, perms) == line


def test_each_line_lands_on_its_workspace():
    everything = ["risk:read", "internal_audit:read", "board:read"]
    for line, key in (("first_line", "my_work"), ("second_line", "dashboard"), ("audit", "assurance"),
                      ("board", "board")):
        assert ws.default_workspace(line, ws.available_workspaces(everything)) == key
    assert ws.default_workspace("first_line", ws.available_workspaces(everything), admin=True) == "dashboard"


def test_a_workspace_needs_its_permission_and_module():
    assert ws.available_workspaces([]) == ["my_work"]
    assert ws.available_workspaces(["internal_audit:read"], modules_off={"internal_audit"}) == ["my_work"]
    # An auditor without internal_audit:read falls back to the dashboard, then My Work.
    assert ws.default_workspace("audit", ["my_work", "dashboard"]) == "dashboard"
    assert ws.default_workspace("board", ["my_work"]) == "my_work"


def test_a_preference_wins_only_while_available():
    avail = ["my_work", "dashboard", "board"]
    assert ws.landing_workspace("board", avail, "my_work") == "my_work"
    assert ws.landing_workspace("board", avail, "assurance") == "board"  # no longer available
    out = ws.workspace_read(role_names=["Board Member"], permissions=["board:read", "governance:read"],
                            preference=None)
    assert out["line"] == "board" and out["landing_href"] == "/board"
    assert [w["key"] for w in out["available"]] == ["my_work", "board"]
    assert ws.workspace_read(role_names=["Admin"], permissions=["settings:manage", "risk:read"],
                             preference="bogus")["preference"] is None


# ================================================================= snapshots ===
def test_a_snapshot_is_due_at_month_end_on_first_run_and_after_a_missed_month():
    assert sn.snapshot_due(date(2026, 9, 30), date(2026, 8, 31))
    assert not sn.snapshot_due(date(2026, 9, 30), date(2026, 9, 30))
    assert not sn.snapshot_due(date(2026, 9, 17), date(2026, 8, 31))
    assert sn.snapshot_due(date(2026, 9, 17), None)
    assert sn.snapshot_due(date(2026, 9, 17), date(2026, 7, 31))  # August's end missed
    assert sn.month_end(date(2026, 2, 3)) == date(2026, 2, 28) and sn.month_end(date(2026, 12, 3)) == date(2026, 12, 31)


def test_quarter_ends_and_the_nearest_snapshot():
    assert sn.quarter_ends(TODAY) == [date(2025, 9, 30), date(2025, 12, 31), date(2026, 3, 31), date(2026, 6, 30)]
    assert sn.quarter_ends(date(2026, 8, 5), 2, fiscal_start_month=8) == [date(2026, 4, 30), date(2026, 7, 31)]
    avail = [date(2026, 6, 28), date(2026, 7, 2), date(2026, 3, 31)]
    assert sn.nearest_date(avail, date(2026, 6, 30)) == date(2026, 6, 28)  # tie → on or before
    assert sn.nearest_date(avail, date(2026, 7, 1)) == date(2026, 7, 2)
    assert sn.nearest_date([date(2026, 1, 1)], date(2026, 6, 30)) is None  # too far
    pts = sn.trend_points(avail, [date(2025, 12, 31), date(2026, 3, 31)])
    assert [p.as_of for p in pts] == [None, date(2026, 3, 31)]


def _overview():
    cat = SimpleNamespace(category_id=uuid.uuid4(), label="Operational", appetite_score=6, tolerance_score=12,
                          risks=4, within_appetite=2, elevated=1, breach=1)
    top = SimpleNamespace(id=uuid.uuid4(), reference="R-1", title="Outage", score=20, severity="critical",
                          appetite_status="breach", owner="Ayesha", status="assessed")
    return SimpleNamespace(
        health=SimpleNamespace(score=64, band="elevated"),
        posture=SimpleNamespace(total_risks=10, within_appetite=5, elevated=3, breach=2, appetite_score=6,
                                tolerance_score=12, by_category=[cat], by_inherent_severity={"high": 3},
                                by_residual_severity={"medium": 3}, top_risks=[top]),
        assurance=SimpleNamespace(total=12, effective=4, partially_effective=2, ineffective=1, not_assessed=3,
                                  not_operating=2, tests_overdue=1, last_test_failed=1),
        compliance=SimpleNamespace(overall_assured_pct=40.0, frameworks=[SimpleNamespace(
            id=uuid.uuid4(), name="ISO 27001", applicable=100, assured=40, unassessed=50, failing=2, unmapped=8,
            gaps=60, compliant_pct=12.5)]),
        kris=SimpleNamespace(red=1, amber=0, red_items=[]),
    )


def test_a_capture_writes_every_headline():
    kri = SimpleNamespace(id=uuid.uuid4(), reference="K-1", name="Failed logins", current_value=9, unit="%",
                          status="red")
    issues = bp.issue_summary([], {}, TODAY)
    rows = sn.rows_from_overview(_overview(), heatmap={"size": 5, "cells": {"5,4": 1}}, kris=[kri], issues=issues)
    keys = {k for k, _d, _v in rows}
    assert set(sn.KEYS) == keys
    by = {(k, d): v for k, d, v in rows}
    assert by[(sn.ASSURANCE, "")]["assured_pct"] == 60.0
    assert by[(sn.KRI_STATUS, "")]["red"] == 1 and by[(sn.KRI_VALUE, str(kri.id))]["value"] == 9.0
    assert by[(sn.APPETITE_TOTAL, "")]["breach"] == 2
    assert by[(sn.TOP_RISKS, "")]["rows"][0]["reference"] == "R-1"


def test_reconstruction_from_history_uses_the_snapshot_at_the_date():
    from app.services.risk_scoring import AppetiteBook, SeverityScale

    r1, r2, r3 = (SimpleNamespace(id=uuid.uuid4()) for _ in range(3))
    history = {
        r1.id: {"status": "assessed", "last_assessed_at": "2026-05-01T00:00:00+00:00", "inherent_likelihood": 5,
                "inherent_impact": 5, "residual_likelihood": 4, "residual_impact": 4, "residual_score": 16},
        r2.id: {"status": "draft", "last_assessed_at": None, "inherent_likelihood": 1, "inherent_impact": 1},
        r3.id: {"status": "assessed", "deleted": True, "inherent_likelihood": 5, "inherent_impact": 5},
    }
    rows = sn.reconstruct_risk_rows([r1, r2, r3], history, AppetiteBook(appetite=6, tolerance=12), SeverityScale(), {})
    total = next(v for k, _d, v in rows if k == sn.APPETITE_TOTAL)
    assert (total["risks"], total["breach"]) == (1, 1) and "basis" in total
    heat = next(v for k, _d, v in rows if k == sn.HEATMAP)
    assert heat["cells"] == {"4,4": 1}

    kri = SimpleNamespace(id=uuid.uuid4(), reference="K", name="K", unit="", direction="higher_is_worse",
                          warning_threshold=3, limit_threshold=5, lower_bound=None, upper_bound=None)
    kr = sn.reconstruct_kri_rows([kri], {kri.id: 4})
    assert kr[0][2]["status"] == "amber" and kr[-1][2]["amber"] == 1

    issues = [SimpleNamespace(created_at=datetime(2026, 5, 1, tzinfo=timezone.utc), closed_date=None,
                              due_date=date(2026, 6, 1), severity="high"),
              SimpleNamespace(created_at=datetime(2026, 5, 1, tzinfo=timezone.utc), closed_date=date(2026, 6, 15),
                              due_date=None, severity="low"),
              SimpleNamespace(created_at=datetime(2026, 7, 5, tzinfo=timezone.utc), closed_date=None,
                              due_date=None, severity="low")]
    _k, _d, v = sn.reconstruct_issue_row(issues, date(2026, 6, 30))
    assert v["open"] == 1 and v["overdue"] == 1 and v["by_severity"]["high"] == 1


def test_trend_series_leave_gaps_where_there_is_no_snapshot():
    rows = [SimpleNamespace(as_of=date(2026, 6, 30), key=sn.APPETITE_TOTAL, dimension="",
                            value={"risks": 8, "within": 5, "elevated": 2, "breach": 1}),
            SimpleNamespace(as_of=date(2026, 6, 30), key=sn.APPETITE_CATEGORY, dimension="default",
                            value={"label": "Default", "breach": 1}),
            SimpleNamespace(as_of=date(2026, 6, 30), key=sn.ASSURANCE, dimension="", value={"assured_pct": 55.0})]
    index = sn.index_rows(rows)
    pts = sn.trend_points([date(2026, 6, 30)], [date(2026, 3, 31), date(2026, 6, 30)])
    series = sn.appetite_series(index, pts)
    assert series[0]["breach"] is None and series[1]["breach"] == 1
    assert sn.category_series(index, pts)["default"]["points"][1]["breach"] == 1
    assert [p["value"] for p in sn.metric_series(index, pts, sn.ASSURANCE, "assured_pct")] == [None, 55.0]


def test_heatmap_counts_the_board_register_at_its_current_cell():
    rows = [SimpleNamespace(status="assessed", last_assessed_at=None, inherent_likelihood=5, inherent_impact=5,
                            residual_likelihood=2, residual_impact=3, residual_score=6),
            SimpleNamespace(status="assessed", last_assessed_at=None, inherent_likelihood=4, inherent_impact=4,
                            residual_likelihood=None, residual_impact=None, residual_score=None),
            SimpleNamespace(status="draft", last_assessed_at=None, inherent_likelihood=5, inherent_impact=5,
                            residual_likelihood=None, residual_impact=None, residual_score=None)]
    assert sn.heatmap_cells(rows, 5) == {"size": 5, "cells": {"2,3": 1, "4,4": 1}}


# ================================================================= assurance ===
def test_findings_follow_up_by_age_and_owner():
    def f(ref, created, due, owner="", status="open", rating="high"):
        return SimpleNamespace(id=uuid.uuid4(), reference=ref, title=ref, rating=rating, status=status,
                               action_owner=owner, engagement_id=uuid.uuid4(), engagement="IT audit",
                               due_date=due, created_at=datetime.combine(created, datetime.min.time(), timezone.utc))
    out = ws.findings_follow_up([
        f("F-1", TODAY - timedelta(days=10), TODAY + timedelta(days=5), "Ayesha"),
        f("F-2", TODAY - timedelta(days=200), TODAY - timedelta(days=100), "Ayesha"),
        f("F-3", TODAY - timedelta(days=60), TODAY - timedelta(days=1)),
        f("F-4", TODAY - timedelta(days=60), TODAY - timedelta(days=1), status="closed"),
    ], TODAY)
    assert out["findings_open"] == 3 and out["findings_overdue"] == 2
    buckets = {b["key"]: (b["count"], b["overdue"]) for b in out["age_buckets"]}
    assert buckets == {"0_30": (1, 0), "31_90": (1, 1), "91_180": (0, 0), "over_180": (1, 1)}
    assert out["by_owner"][0]["owner"] == "Ayesha" and out["by_owner"][0]["oldest_days"] == 200
    assert [r["reference"] for r in out["overdue_findings"]] == ["F-2", "F-3"]
    assert out["overdue_findings"][1]["owner"] == ""


def test_the_three_lines_map_shows_coverage_and_gaps():
    cat_a, cat_b = uuid.uuid4(), uuid.uuid4()
    r1, r2, r3 = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    c1, c2 = uuid.uuid4(), uuid.uuid4()
    recent, stale = TODAY - timedelta(days=30), TODAY - timedelta(days=500)
    inputs = ws.MapInputs(
        categories={cat_a: "Cyber", cat_b: "Credit"},
        risk_category={r1: cat_a, r2: cat_b, r3: None},
        risk_controls=[(r1, c1), (r1, c2)],
        control_attested={c1: recent, c2: recent},
        risk_attested={},
        rcsa_completed={},
        control_tested={c1: recent, c2: stale},
        control_clauses={c1: 3},
        engagements=[(uuid.uuid4(), recent, frozenset({r1}), frozenset(), ""),
                     (uuid.uuid4(), stale - timedelta(days=900), frozenset(), frozenset(), "Credit")],
    )
    rows = {r["label"]: r for r in ws.assurance_map(inputs, TODAY)}
    cyber, credit, none = rows["Cyber"], rows["Credit"], rows["No category"]
    assert (cyber["risks"], cyber["controls"]) == (1, 2)
    assert cyber["first_line"]["state"] == "covered" and cyber["first_line"]["pct"] == 100.0
    assert cyber["second_line"]["state"] == "partial" and cyber["second_line"]["last"] == recent
    assert "3 assessed compliance clause" in cyber["second_line"]["detail"]
    assert cyber["third_line"]["state"] == "covered"
    assert cyber["gaps"] == ["Second line: controls not independently tested in the last 12 months"]
    assert credit["second_line"]["state"] == "not_applicable"
    assert credit["third_line"]["state"] == "partial"  # audited by unit category, but over 3 years ago
    assert "No controls linked to the category's risks" in credit["gaps"]
    assert none["key"] == ws.UNCATEGORISED and none["third_line"]["state"] == "none"


def test_engagement_end_dates():
    assert ws.engagement_end("fieldwork", None, None, date(2026, 1, 1)) is None
    assert ws.engagement_end("closed", None, None, date(2026, 1, 1)) == date(2026, 1, 1)
    assert ws.engagement_end("reporting", date(2026, 2, 1), date(2026, 3, 1), None) == date(2026, 3, 1)


def test_decisions_due_put_overdue_and_mine_first():
    me = uuid.uuid4()

    def d(ref, due, owner_id=None, status="open"):
        return SimpleNamespace(id=uuid.uuid4(), reference=ref, description=ref, decision_type="action",
                               status=status, owner="", owner_id=owner_id, due_date=due, committee="BRMC",
                               meeting="Q3", meeting_date=None)
    rows = ws.decision_rows([d("A", TODAY + timedelta(days=3)), d("B", TODAY + timedelta(days=5), me),
                             d("C", TODAY - timedelta(days=2)), d("D", TODAY + timedelta(days=90)),
                             d("E", TODAY - timedelta(days=9), status="done")], user_id=me, today=TODAY)
    assert [r["reference"] for r in rows] == ["C", "B", "A"]
    assert rows[1]["mine"] and rows[0]["overdue"]


def test_the_latest_ready_pack_decides_a_meetings_pack_state():
    m = uuid.uuid4()
    t = datetime(2026, 9, 1, tzinfo=timezone.utc)
    packs = [SimpleNamespace(meeting_id=m, status="ready", review_state="released", at=t),
             SimpleNamespace(meeting_id=m, status="ready", review_state="draft", at=t + timedelta(days=1)),
             SimpleNamespace(meeting_id=m, status="failed", review_state="draft", at=t + timedelta(days=2))]
    assert ws.pack_state_for_meetings(packs) == {m: "draft"}


# =============================================================== board packs ===
MAKER, CHECKER = uuid.uuid4(), uuid.uuid4()


def refusal(action, state, actor, required=True, status="ready"):
    return bp.lifecycle_refusal(action=action, review_state=state, status=status, actor_id=actor,
                                contributor_ids=[str(MAKER)], dual_control=required)


def test_sign_off_follows_draft_reviewed_released_with_four_eyes():
    assert refusal("review", "draft", CHECKER) is None
    assert refusal("review", "draft", MAKER).startswith("Segregation of duties")
    assert refusal("review", "draft", MAKER, required=False) is None
    assert "Only a draft" in refusal("review", "reviewed", CHECKER)
    assert refusal("release", "draft", CHECKER) == "A pack must be reviewed before it is released."
    assert refusal("release", "reviewed", CHECKER) is None
    assert refusal("release", "reviewed", MAKER).startswith("Segregation of duties")
    assert refusal("release", "released", CHECKER) == "This pack is already released."
    assert refusal("return", "reviewed", MAKER) is None and refusal("return", "draft", MAKER)
    assert refusal("commentary", "released", MAKER).startswith("A released pack is final")
    assert refusal("commentary", "reviewed", MAKER) is None
    assert refusal("review", "draft", CHECKER, status="failed")


def test_commentary_is_for_the_packs_own_sections():
    assert bp.clean_commentary(["summary", "kris"], {"summary": "  Risk rose.  ", "kris": " "}) == {"summary": "Risk rose."}
    with pytest.raises(bp.PackError, match="no 'issues' section"):
        bp.clean_commentary(["summary"], {"issues": "x"})
    with pytest.raises(bp.PackError, match="longer than"):
        bp.clean_commentary(["summary"], {"summary": "x" * (bp.COMMENTARY_MAX + 1)})


def test_a_committee_keeps_its_section_order():
    assert bp.normalise_sections(["kris", "summary", "kris"], keep_order=True) == ["kris", "summary"]
    assert bp.normalise_sections(["kris", "summary"]) == ["summary", "kris"]


def test_distribution_goes_to_active_members_once():
    uid = uuid.uuid4()
    rows = [SimpleNamespace(user_id=uid, full_name="Zara", email="z@bank.pk", is_active=True),
            SimpleNamespace(user_id=uid, full_name="Zara", email="z@bank.pk", is_active=True),
            SimpleNamespace(user_id=uuid.uuid4(), full_name="", email="gone@bank.pk", is_active=False),
            SimpleNamespace(user_id=uuid.uuid4(), full_name="Ali", email="a@bank.pk", is_active=True)]
    assert [r["name"] for r in bp.distribution_list(rows)] == ["Ali", "Zara"]


def _pack(**sections):
    return {
        "title": "Board pack: BRMC", "sections": list(sections),
        "cover": {"organisation": "Demo Bank", "committee": "BRMC", "period_start": date(2026, 4, 1),
                  "period_end": date(2026, 6, 30), "as_of": TODAY,
                  "generated_at": datetime(2026, 9, 17, 9, tzinfo=timezone.utc), "generated_by": "Jane",
                  "timezone": "Asia/Karachi", "tz": KHI, "date_format": "DD/MM/YYYY"},
        **sections,
    }


def test_stored_figures_round_trip():
    pack = _pack(kris=bp.kri_summary([]))
    pack["basis"] = {"position": "snapshot", "snapshot_as_of": date(2026, 6, 30)}
    import json

    back = bp.decode_content(json.loads(json.dumps(bp.encode_content(pack))))
    assert back["cover"]["period_end"] == date(2026, 6, 30)
    assert back["cover"]["generated_at"] == pack["cover"]["generated_at"]
    assert back["cover"]["tz"] == KHI and back["basis"]["snapshot_as_of"] == date(2026, 6, 30)


def test_a_past_period_uses_the_snapshot_and_says_so():
    kri_id = str(uuid.uuid4())
    figures = {
        sn.HEALTH: {"": {"score": 55, "band": "elevated"}},
        sn.APPETITE_TOTAL: {"": {"risks": 7, "within": 4, "elevated": 2, "breach": 1, "appetite": 6, "tolerance": 12}},
        sn.KRI_STATUS: {"": {"green": 1, "amber": 0, "red": 1, "no_data": 0}},
        sn.KRI_VALUE: {kri_id: {"reference": "K-1", "name": "Logins", "status": "red", "value": 9.0, "unit": "%"}},
    }
    pack = _pack(summary={"score": 70, "band": "good", "headlines": [], "actions": [{"label": "x", "count": 1}],
                          "components": []},
                 appetite={"rows": [], "total": 10}, assurance={"effective": 9, "failed": []},
                 kris=bp.kri_summary([]))
    bp.apply_snapshot(pack, figures, date(2026, 6, 30))
    pack["basis"] = {"position": "snapshot", "snapshot_as_of": date(2026, 6, 30)}
    assert pack["summary"]["score"] == 55 and pack["summary"]["actions"] == []
    assert pack["appetite"]["breach"] == 1 and pack["appetite"]["rows"][0]["risks"] == 7
    assert pack["assurance"]["effective"] == 9  # no assurance snapshot: stays, and says so
    assert "as at the day the pack was generated" in pack["position_notes"]["assurance"][0]
    assert pack["kris"]["red"] == 1 and pack["kris"]["rows"][0]["id"] == kri_id
    facts = dict(bp.cover_facts(pack))
    assert "snapshot nearest the period end" in facts["Position as at"]
    views = {v.key: v for v in bp.section_views(pack)}
    assert views["appetite"].notes[0].startswith("Position figures are from the period snapshot of 30/06/2026")


def test_commentary_charts_and_branding_render():
    kri = SimpleNamespace(id=uuid.uuid4(), reference="K-1", name="Failed logins", status="red", current_value=9,
                          unit="%", warning_threshold=3, limit_threshold=5, lower_bound=None, upper_bound=None,
                          direction="higher_is_worse", owner="", owner_id=None, last_measured_date=TODAY)
    kris = bp.kri_summary([kri])
    kris["trend"] = {str(kri.id): [{"as_of": date(2026, 7, 1), "value": 4.0}, {"as_of": date(2026, 8, 1), "value": 9.0}]}
    appetite = {"total": 3, "within": 1, "elevated": 1, "breach": 1, "appetite": 6, "tolerance": 12,
                "rows": [{"label": "All", "appetite": 6, "tolerance": 12, "risks": 3, "within": 1, "elevated": 1, "breach": 1}],
                "heatmap": {"size": 5, "cells": {"5,5": 1, "2,2": 2}, "bands": {"5,5": "critical", "2,2": "low"}},
                "trend": [{"date": date(2026, 3, 31), "within": None, "elevated": None, "breach": None},
                          {"date": date(2026, 6, 30), "within": 1, "elevated": 1, "breach": 1}]}
    pack = _pack(appetite=appetite, kris=kris)
    pack["commentary"] = {"appetite": "Cyber exposure <rose> after the outage."}
    pack["review_state"] = bp.DRAFT
    pack["cover"]["branding"] = {"cover_title": "Board Risk Committee", "primary_colour": "#0a7a55",
                                 "classification": "Strictly confidential"}
    views = {v.key: v for v in bp.section_views(pack)}
    assert views["appetite"].commentary.startswith("Cyber") and {c.kind for c in views["appetite"].charts} == {
        "heatmap", "appetite_trend"}
    assert views["kris"].charts[0].kind == "kri_trend"
    assert bp.to_pdf(pack).startswith(b"%PDF")
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(bp.to_xlsx(pack)))
    values = [c.value for row in wb[bp.SECTION_TITLES["appetite"]].iter_rows() for c in row]
    assert "Commentary" in values and "No snapshot" in values
    assert bp.brand_colour({"primary_colour": "red"}) == bp.brand_colour(None)
    assert bp.brand_colour({"primary_colour": "#0a7a55"}) == "#0a7a55"


class _Nested:
    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False


class PackDB:
    def __init__(self):
        self.added = []

    def begin_nested(self):
        return _Nested()

    def add(self, obj):
        self.added.append(obj)

    async def flush(self):
        pass


def _org():
    return bp.OrgContext(tenant_id=uuid.uuid4(), name="Demo Bank", tz=KHI, timezone="Asia/Karachi",
                         date_format="DD/MM/YYYY", fiscal_start_month=1, today=TODAY,
                         now=datetime(2026, 9, 17, 9, tzinfo=timezone.utc))


async def test_generation_starts_a_draft_with_the_committees_sections_and_keeps_the_figures(monkeypatch):
    from app.services import audit

    async def fake_build(db, org, **kw):
        return {"title": kw["title"], "sections": kw["sections"], "cover": {"period_end": kw["period_end"]},
                "basis": {"position": "live"}}

    async def fake_store(db, tenant_id, **kw):
        return SimpleNamespace(id=uuid.uuid4(), storage_key="k", filename=kw["filename"])

    async def fake_record(db, **kw):
        pass

    monkeypatch.setattr(bp, "build_pack", fake_build)
    monkeypatch.setattr(bp, "to_pdf", lambda data: b"%PDF")
    monkeypatch.setattr(bp, "to_xlsx", lambda data: b"PK")
    monkeypatch.setattr(bp, "store_file", fake_store)
    monkeypatch.setattr(audit, "record", fake_record)
    user = SimpleNamespace(id=MAKER, email="m@bank.pk", full_name="Maker", tenant_id=uuid.uuid4())
    committee = SimpleNamespace(id=uuid.uuid4(), name="BRMC", reference="CMT-1", board_pack_sections=["kris", "summary"])
    row = await bp.generate(PackDB(), _org(), actor=user, committee=committee)
    assert row.sections == ["kris", "summary"] and row.review_state == bp.DRAFT
    assert row.contributor_ids == [str(MAKER)]
    assert row.content["cover"]["period_end"] == {"$d": "2026-09-17"}


async def test_review_and_release_enforce_four_eyes(monkeypatch):
    from fastapi import HTTPException  # noqa: F401

    from app.services import audit, dual_control, email

    async def required(db, module, action, amount=None):
        assert (module, action) == ("board_pack", "release")
        return True, None

    async def fake_record(db, **kw):
        pass

    async def recipients(db, committee_id):
        return [{"user_id": str(uuid.uuid4()), "name": "Zara", "email": "z@bank.pk", "emailed": False}]

    sent = []

    async def fake_send(to, subject, html, text=None):
        sent.append((to, subject))
        return True

    rendered = []

    async def fake_render(db, org, pack, *, uploader):
        rendered.append(pack.review_state)

    monkeypatch.setattr(dual_control, "dual_control_required", required)
    monkeypatch.setattr(audit, "record", fake_record)
    monkeypatch.setattr(bp, "committee_recipients", recipients)
    monkeypatch.setattr(bp, "render_stored", fake_render)
    monkeypatch.setattr(email, "send_email", fake_send)
    pack = SimpleNamespace(id=uuid.uuid4(), title="Q3 pack", status="ready", review_state="draft",
                           contributor_ids=[str(MAKER)], committee_id=uuid.uuid4(), period_start=date(2026, 7, 1),
                           period_end=date(2026, 9, 17), content={"x": 1}, reviewed_by_id=None, reviewed_at=None,
                           released_by_id=None, released_at=None, distribution=[])
    maker = SimpleNamespace(id=MAKER, email="m@bank.pk", tenant_id=uuid.uuid4())
    checker = SimpleNamespace(id=CHECKER, email="c@bank.pk", tenant_id=uuid.uuid4())
    db = PackDB()
    with pytest.raises(bp.PackError, match="Segregation"):
        await bp.review_pack(db, pack, maker)
    await bp.review_pack(db, pack, checker)
    assert pack.review_state == "reviewed" and pack.reviewed_by_id == CHECKER
    with pytest.raises(bp.PackError, match="Segregation"):
        await bp.release_pack(db, _org(), pack, maker)
    await bp.release_pack(db, _org(), pack, checker)
    assert pack.review_state == "released" and rendered == ["released"]
    assert pack.distribution[0]["emailed"] is True and sent[0][0] == ["z@bank.pk"]
    note = [o for o in db.added if o.__class__.__name__ == "Notification"][0]
    assert note.link == f"/board?pack={pack.id}" and note.dedup_key.startswith("event:")


async def test_commentary_on_a_reviewed_pack_returns_it_to_draft(monkeypatch):
    from app.services import audit

    async def fake_record(db, **kw):
        fake_record.summary = kw["summary"]

    async def fake_render(db, org, pack, *, uploader):
        pass

    monkeypatch.setattr(audit, "record", fake_record)
    monkeypatch.setattr(bp, "render_stored", fake_render)
    pack = SimpleNamespace(id=uuid.uuid4(), title="Q3", status="ready", review_state="reviewed", sections=["summary"],
                           commentary={}, contributor_ids=[str(MAKER)], reviewed_by_id=CHECKER, reviewed_at=TODAY)
    writer = SimpleNamespace(id=CHECKER, email="c@bank.pk")
    await bp.set_commentary(PackDB(), _org(), pack, writer, {"summary": "Appetite held."})
    assert pack.review_state == "draft" and pack.reviewed_by_id is None
    assert str(CHECKER) in pack.contributor_ids  # the writer can no longer review it
    assert "returned to draft" in fake_record.summary


# ====================================================== SQL compiles (no rows) ===
class CompilingDB:
    def __init__(self, scalar=None):
        self.sql: list[str] = []
        self._scalar = scalar

    def _compile(self, stmt):
        self.sql.append(str(stmt.compile(dialect=postgresql.dialect())))

    async def execute(self, stmt, *a, **kw):
        self._compile(stmt)
        return SimpleNamespace(all=lambda: [])

    async def scalars(self, stmt, *a, **kw):
        self._compile(stmt)
        return SimpleNamespace(all=lambda: [])

    async def scalar(self, stmt, *a, **kw):
        self._compile(stmt)
        return self._scalar


async def test_the_assurance_and_snapshot_loaders_compile():
    db = CompilingDB()
    inputs = await ws.load_map_inputs(db, TODAY)
    assert ws.assurance_map(inputs, TODAY) == []
    assert await sn.available_dates(db, since=TODAY) == []
    assert await sn.live_heatmap(db) == {"size": 5, "cells": {}}
    assert await bp.snapshot_for(db, date(2026, 6, 30)) is None
    assert await bp.load_branding(db) == {"cover_title": "", "primary_colour": "", "classification": "Confidential",
                                          "logo_path": None}
    joined = "\n".join(db.sql)
    assert "audit_finding_risks" in joined and "requirement_controls" in joined


async def test_a_past_period_pack_with_no_snapshot_says_so(monkeypatch):
    async def fake_overview(db, tenant_id, days, viewer=None):
        from tests.test_board_pack import _overview as ov

        return ov()

    monkeypatch.setattr(bp, "dashboard_overview", fake_overview)
    pack = await bp.build_pack(CompilingDB(), _org(), period_start=date(2026, 4, 1), period_end=date(2026, 6, 30),
                               sections=["appetite", "kris"], title="Q2")
    assert pack["basis"] == {"position": "live", "reason": "no_snapshot"}
    assert "heatmap" in pack["appetite"] and pack["appetite"]["trend"][-1]["date"] == date(2026, 6, 30)
    assert "no period snapshot" in dict(bp.cover_facts(pack))["Position as at"]


async def test_the_board_and_assurance_homes_build_from_empty_data(monkeypatch):
    from app.services import modules

    async def fake_overview(db, tenant_id, days, viewer=None):
        return _overview()

    monkeypatch.setattr(bp, "dashboard_overview", fake_overview)
    monkeypatch.setattr(modules, "is_enabled", lambda key: True)
    user = SimpleNamespace(id=ME, tenant_id=uuid.uuid4(), permission_codes=["board:read"])
    db = CompilingDB()
    home = await ws.board_home(db, user, _org())
    from app.schemas.workspaces import AssuranceHome, BoardHome

    board = BoardHome(**home)
    assert board.appetite.categories[0].trend[-1].breach == 1 and len(board.appetite.trend) == 5
    assert board.assurance.value == 60.0 and board.governance_enabled and board.decisions == []
    assert board.top_risks[0].movement == "unknown"
    assurance = AssuranceHome(**await ws.assurance_home(db, user, _org()))
    assert assurance.internal_audit_enabled and assurance.engagements == [] and not assurance.document_requests_supported
    joined = "\n".join(db.sql)
    assert "meeting_decisions" in joined and "audit_procedures" not in joined  # no engagements, no procedure query


# =================================================================== My Work ===
ME = uuid.uuid4()


def _ctx(perms=("assessment:read", "internal_audit:read", "review:read", "regchange:read", "exception:read",
                "declaration:read")):
    directory = Directory(users={ME: DirectoryUser(id=ME, email="ayesha@bank.pk", full_name="Ayesha Siddiqui")})
    return mw.Ctx(user_id=ME, email="ayesha@bank.pk", permissions=set(perms), role_names=set(), role_ids=set(),
                  today=TODAY, horizon=TODAY + timedelta(days=mw.HORIZON_DAYS), directory=directory)


def test_free_text_people_are_matched_exactly():
    ctx = _ctx()
    assert mw.names_me(ctx, "Ayesha Siddiqui")
    assert mw.names_me(ctx, "Bilal Khan, ayesha@bank.pk")
    assert not mw.names_me(ctx, "Ayesha", "")


def test_the_new_kinds_are_registered_in_order():
    new = ["assessment_review", "finding_follow_up", "engagement_task", "audit_remediation", "access_review",
           "assessment_chase", "regulatory_change", "regulatory_return", "incident_report", "exception_expiry",
           "declaration"]
    kinds = [k for k, _l, _h in mw.KINDS]
    assert kinds[-len(new):] == new and set(mw.BUILDERS) == set(kinds)


class RowsDB(CompilingDB):
    """Answers each statement with the rows registered for the first table it names."""

    def __init__(self, rows):
        super().__init__()
        self.rows = rows

    async def execute(self, stmt, *a, **kw):
        self._compile(stmt)
        sql = self.sql[-1]
        for table, rows in self.rows.items():
            if f"FROM {table}" in sql:
                return SimpleNamespace(all=lambda rows=rows: rows)
        return SimpleNamespace(all=lambda: [])


async def test_every_new_kind_builds_valid_sql():
    db = CompilingDB()
    ctx = _ctx()
    for kind in ("assessment_review", "finding_follow_up", "engagement_task", "audit_remediation", "access_review",
                 "assessment_chase", "regulatory_change", "regulatory_return", "incident_report",
                 "exception_expiry", "declaration"):
        assert await mw.BUILDERS[kind](db, ctx) == []
    assert len(db.sql) >= 11


async def test_new_kinds_pick_my_items():
    eng = uuid.uuid4()
    rows = {
        "access_reviews": [SimpleNamespace(id=uuid.uuid4(), reference="AR-1", name="Core banking users",
                                           reviewer="Ayesha Siddiqui", system_name="T24", due_date=TODAY, status="in_progress"),
                           SimpleNamespace(id=uuid.uuid4(), reference="AR-2", name="Other", reviewer="Bilal",
                                           system_name="", due_date=TODAY, status="in_progress")],
        "exceptions": [SimpleNamespace(id=uuid.uuid4(), reference="EX-1", title="Legacy TLS", expires_at=TODAY + timedelta(days=20),
                                       requested_by=None, approver_id=ME, business_owner=""),
                       SimpleNamespace(id=uuid.uuid4(), reference="EX-2", title="Not mine", expires_at=TODAY,
                                       requested_by=None, approver_id=None, business_owner="Bilal")],
        "audit_findings": [SimpleNamespace(id=uuid.uuid4(), reference="F-1", title="Weak MFA", rating="high",
                                           action_owner="ayesha@bank.pk", due_date=TODAY, engagement_id=eng,
                                           engagement="IT general controls")],
    }
    ctx = _ctx()
    reviews = await mw.my_access_reviews(RowsDB(rows), ctx)
    assert [i.reference for i in reviews] == ["AR-1"]
    expiring = await mw.my_expiring_exceptions(RowsDB(rows), ctx)
    assert [i.reference for i in expiring] == ["EX-1"] and "you approved it" in expiring[0].subtitle
    fix = await mw.my_audit_remediation(RowsDB(rows), ctx)
    assert [i.reference for i in fix] == ["F-1"] and fix[0].link == f"/internal-audit?id={eng}"
    ctx.modules_off = {"internal_audit"}
    assert await mw.my_audit_remediation(RowsDB(rows), ctx) == []


# ===================================================================== roles ===
def test_the_board_member_role_arrives_once():
    from app.core.permissions import DEFAULT_ROLES, PERMISSION_CATALOG
    from app.db.provisioning import IMPLIED_ON_INTRODUCTION, implied_grants, roles_to_introduce

    assert "board:read" in PERMISSION_CATALOG and "boardpack:release" in PERMISSION_CATALOG
    assert DEFAULT_ROLES["Board Member"][1] == ["board:read", "governance:read"]
    assert "board:read" in DEFAULT_ROLES["Auditor"][1] and "board:read" in DEFAULT_ROLES["Viewer"][1]
    assert roles_to_introduce({"board:read"}, {"Admin", "Viewer"}) == ["Board Member"]
    assert roles_to_introduce({"board:read"}, {"board member"}) == []
    assert roles_to_introduce(set(), {"Admin"}) == []
    assert implied_grants({"boardpack:release"}, {"governance:write"}) == {"boardpack:release"}
    assert IMPLIED_ON_INTRODUCTION["board:read"] == "governance:read"


def test_schema_patch_is_idempotent_and_registered():
    from app.core.database import Base
    from app.db.phase4 import all_tables
    from app.db.phase4.workspaces import TABLES, ddl_statements
    from app.db.rls import TENANT_SCOPED_TABLES

    import app.models  # noqa: F401

    assert set(TABLES) <= set(all_tables()) and all(t in Base.metadata.tables for t in TABLES)
    assert set(TABLES) <= set(TENANT_SCOPED_TABLES)
    for statement in ddl_statements():
        assert "IF NOT EXISTS" in statement or "SET DEFAULT" in statement

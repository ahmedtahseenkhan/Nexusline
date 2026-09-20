"""Board packs (phase 3, plan §3.4).

The pack is the version of the numbers a committee saw, so these pin the rules that
decide them: the default period (the fiscal quarter to date), the section choice, the
top-risk trend and what it rests on, the period movement, the issue / incident / KRI /
third-party summaries, what the scheduler considers due, that the PDF and spreadsheet
print the same views, and that a failed generation is kept as a failed pack with no
files. Pure — no database (the orchestration runs against a fake session).
"""
import io
import uuid
from datetime import date, datetime, timezone
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest

from app.models.enums import KriStatus
from app.services import board_pack as bp

KHI = ZoneInfo("Asia/Karachi")
UTC = timezone.utc


def at(y, m, d, h=0, tz=UTC):
    return datetime(y, m, d, h, tzinfo=tz)


# ---------------------------------------------------------------- period & sections ---
@pytest.mark.parametrize("day,fy,expected", [
    (date(2026, 9, 12), 1, date(2026, 7, 1)),
    (date(2026, 1, 1), 1, date(2026, 1, 1)),
    (date(2026, 12, 31), 1, date(2026, 10, 1)),
    (date(2026, 9, 12), 7, date(2026, 7, 1)),   # July fiscal year: Q1 is Jul-Sep
    (date(2026, 6, 30), 7, date(2026, 4, 1)),   # ... and Q4 Apr-Jun
    (date(2026, 1, 15), 2, date(2025, 11, 1)),  # February fiscal year wraps the year
    (date(2026, 5, 5), 4, date(2026, 4, 1)),
])
def test_fiscal_quarter_start(day, fy, expected):
    assert bp.fiscal_quarter_start(day, fy) == expected


def test_default_period_is_the_quarter_to_date():
    assert bp.resolve_period(None, None, today=date(2026, 9, 12)) == (date(2026, 7, 1), date(2026, 9, 12))


def test_an_end_alone_runs_from_the_start_of_its_quarter():
    assert bp.resolve_period(None, date(2026, 6, 30), today=date(2026, 9, 12)) == (date(2026, 4, 1), date(2026, 6, 30))


def test_a_period_ending_in_the_future_is_refused():
    with pytest.raises(bp.PackError, match="future"):
        bp.resolve_period(None, date(2026, 9, 13), today=date(2026, 9, 12))


def test_a_period_starting_after_it_ends_is_refused():
    with pytest.raises(bp.PackError, match="on or before"):
        bp.resolve_period(date(2026, 9, 1), date(2026, 8, 1), today=date(2026, 9, 12))


def test_sections_default_to_all_in_pack_order():
    assert bp.normalise_sections(None) == list(bp.SECTION_KEYS)
    assert bp.normalise_sections(["kris", "summary", "kris"]) == ["summary", "kris"]


def test_unknown_or_no_sections_are_refused():
    with pytest.raises(bp.PackError, match="Unknown section"):
        bp.normalise_sections(["summary", "gossip"])
    with pytest.raises(bp.PackError, match="at least one"):
        bp.normalise_sections([])


@pytest.mark.parametrize("fmt,expected", [
    ("DD/MM/YYYY", "05/09/2026"), ("MM/DD/YYYY", "09/05/2026"),
    ("YYYY-MM-DD", "2026-09-05"), ("DD MMM YYYY", "05 Sep 2026"),
])
def test_dates_follow_the_organisation_format(fmt, expected):
    assert bp.format_date(date(2026, 9, 5), fmt) == expected


def test_times_are_shown_in_the_organisation_timezone():
    assert bp.format_datetime(at(2026, 9, 12, 9), "DD/MM/YYYY", KHI) == "12/09/2026 14:00"
    assert bp.format_date(None) == "—"


def test_pack_title_names_the_committee_and_meeting():
    kw = dict(start=date(2026, 7, 1), end=date(2026, 9, 12), date_format="DD/MM/YYYY")
    assert bp.pack_title(committee="BRMC", meeting="Q3 sitting", **kw) == "Board pack: BRMC — Q3 sitting"
    assert bp.pack_title(committee="BRMC", meeting=None, **kw) == "Board pack: BRMC, 01/07/2026 to 12/09/2026"
    assert bp.pack_title(committee=None, meeting=None, **kw) == "Board pack, 01/07/2026 to 12/09/2026"


def test_period_instants_cover_whole_days_in_the_organisation_zone():
    start, end = bp.period_instants(date(2026, 7, 1), date(2026, 9, 12), KHI)
    assert start == datetime(2026, 7, 1, tzinfo=KHI)
    assert end == datetime(2026, 9, 13, tzinfo=KHI)
    assert bp.in_range(datetime(2026, 9, 12, 23, 59, tzinfo=KHI), start, end)
    assert not bp.in_range(datetime(2026, 9, 13, 0, 0, tzinfo=KHI), start, end)


# ------------------------------------------------------------------ scheduler rules ---
def test_a_meeting_is_due_its_pack_inside_the_window_only():
    today = date(2026, 9, 12)
    assert bp.meeting_due(date(2026, 9, 19), 7, today)
    assert bp.meeting_due(date(2026, 9, 12), 7, today)        # the day itself
    assert not bp.meeting_due(date(2026, 9, 20), 7, today)    # too early
    assert not bp.meeting_due(date(2026, 9, 11), 7, today)    # already held
    assert not bp.meeting_due(None, 7, today)
    assert not bp.meeting_due(date(2026, 9, 19), None, today)
    assert not bp.meeting_due(date(2026, 9, 19), 0, today)


def test_only_a_pack_generated_inside_the_window_counts():
    meeting = date(2026, 9, 19)
    assert bp.pack_is_current(date(2026, 9, 12), meeting, 7)
    assert bp.pack_is_current(date(2026, 9, 18), meeting, 7)
    assert not bp.pack_is_current(date(2026, 8, 1), meeting, 7)  # a draft weeks earlier is stale
    assert not bp.pack_is_current(None, meeting, 7)


def test_lock_keys_are_stable_signed_64_bit():
    mid = uuid.uuid4()
    key = bp._lock_key(mid)
    assert key == bp._lock_key(mid)
    assert -(2 ** 63) <= key < 2 ** 63


# ---------------------------------------------------------------------- risk trend ---
START = at(2026, 7, 1)


def test_a_risk_created_in_the_period_is_new():
    t = bp.risk_trend(current=12, created_at=at(2026, 8, 1), last_assessed_at=at(2026, 8, 1), period_start_at=START)
    assert (t.direction, t.basis) == ("new", bp.TREND_NEW)
    assert bp.trend_text(t) == "New in period"


def test_a_risk_not_rescored_since_the_start_is_unchanged_for_certain():
    t = bp.risk_trend(current=15, created_at=at(2026, 1, 1), last_assessed_at=at(2026, 3, 1), period_start_at=START,
                      snapshot={"inherent_likelihood": 5, "inherent_impact": 5}, snapshot_at=at(2026, 2, 1))
    assert (t.direction, t.start_score, t.basis) == ("same", 15, bp.TREND_UNCHANGED)


@pytest.mark.parametrize("current,direction", [(20, "up"), (6, "down"), (12, "same")])
def test_a_rescored_risk_compares_with_its_history(current, direction):
    snap = {"inherent_likelihood": 4, "inherent_impact": 4, "residual_likelihood": 3, "residual_impact": 4}
    t = bp.risk_trend(current=current, created_at=at(2026, 1, 1), last_assessed_at=at(2026, 8, 1),
                      period_start_at=START, snapshot=snap, snapshot_at=at(2026, 5, 1))
    assert (t.direction, t.start_score, t.basis) == (direction, 12, bp.TREND_HISTORY)


def test_a_residual_assessed_after_the_last_snapshot_is_the_starting_score():
    snap = {"inherent_likelihood": 5, "inherent_impact": 4}  # 20, no residual yet
    t = bp.risk_trend(current=9, created_at=at(2026, 1, 1), last_assessed_at=at(2026, 8, 1), period_start_at=START,
                      snapshot=snap, snapshot_at=at(2026, 2, 1),
                      assessment={"residual_likelihood": 2, "residual_impact": 4}, assessment_at=at(2026, 3, 1))
    assert (t.direction, t.start_score) == ("up", 8)
    # An assessment older than the snapshot is already in it.
    t = bp.risk_trend(current=9, created_at=at(2026, 1, 1), last_assessed_at=at(2026, 8, 1), period_start_at=START,
                      snapshot=snap, snapshot_at=at(2026, 4, 1),
                      assessment={"residual_likelihood": 2, "residual_impact": 4}, assessment_at=at(2026, 3, 1))
    assert t.start_score == 20


def test_no_history_before_the_period_is_said_not_guessed():
    t = bp.risk_trend(current=9, created_at=at(2026, 1, 1), last_assessed_at=at(2026, 8, 1), period_start_at=START)
    assert (t.direction, t.start_score, t.basis) == ("unknown", None, bp.TREND_UNKNOWN)
    assert bp.trend_text(t) == "No earlier score"


def test_trend_text():
    assert bp.trend_text(bp.Trend("up", 8, bp.TREND_HISTORY)) == "Up from 8"
    assert bp.trend_text(bp.Trend("down", 20, bp.TREND_HISTORY)) == "Down from 20"
    assert bp.trend_text(bp.Trend("same", 12, bp.TREND_UNCHANGED)) == "Unchanged"


def test_score_from_snapshot_prefers_the_residual():
    assert bp.score_from_snapshot({"inherent_likelihood": 4, "inherent_impact": 5}) == 20
    assert bp.score_from_snapshot({"inherent_likelihood": 4, "inherent_impact": 5,
                                   "residual_likelihood": 2, "residual_impact": 3}) == 6
    assert bp.score_from_snapshot({}) is None


def test_risk_movement_in_the_period():
    end = at(2026, 9, 13)
    r = lambda ref, **kw: SimpleNamespace(reference=ref, status=kw.get("status", "assessed"),  # noqa: E731
                                          created_at=kw.get("created", at(2026, 1, 1)),
                                          updated_at=kw.get("updated", at(2026, 1, 1)),
                                          last_assessed_at=kw.get("assessed"))
    rows = [
        r("R-1", created=at(2026, 8, 1), assessed=at(2026, 8, 2)),        # new (first scoring isn't a re-score)
        r("R-2", status="closed", updated=at(2026, 8, 5)),                # closed in period
        r("R-3", assessed=at(2026, 7, 15)),                               # re-scored
        r("R-4", status="closed", updated=at(2026, 6, 30)),               # closed before the period
        r("R-5", assessed=at(2026, 6, 1)),                                # re-scored before the period
    ]
    moved = bp.classify_risk_movement(rows, START, end)
    assert [x.reference for x in moved["new"]] == ["R-1"]
    assert [x.reference for x in moved["closed"]] == ["R-2"]
    assert [x.reference for x in moved["rescored"]] == ["R-3"]


# ------------------------------------------------------------------ period sections ---
def _issue(ref, sev, due=None, owner="", owner_id=None, reg=False):
    return SimpleNamespace(id=uuid.uuid4(), reference=ref, title=f"Issue {ref}", severity=sev, due_date=due,
                           owner=owner, owner_id=owner_id, regulator_related=reg)


def test_issue_summary_by_severity_overdue_and_moved():
    today = date(2026, 9, 12)
    owner = uuid.uuid4()
    a = _issue("ISS-1", "high", date(2026, 9, 1), owner_id=owner)
    b = _issue("ISS-2", "critical", date(2026, 9, 10), owner="Legacy Name")
    c = _issue("ISS-3", "high", date(2026, 10, 1), reg=True)
    d = _issue("ISS-4", "low")
    s = bp.issue_summary([a, b, c, d], {a.id: 2, c.id: 1}, today, {owner: "Ali Khan"})
    assert (s["open"], s["overdue"], s["date_moved"], s["regulator_related"]) == (4, 2, 2, 1)
    high = next(x for x in s["by_severity"] if x["severity"] == "high")
    assert high == {"severity": "high", "open": 2, "overdue": 1, "date_moved": 2, "moves": 3}
    assert [x["severity"] for x in s["by_severity"]] == ["critical", "high", "medium", "low"]
    # Worst severity first; owner from the user list, else the legacy text.
    assert [(x["reference"], x["owner"], x["days_overdue"]) for x in s["overdue_rows"]] == [
        ("ISS-2", "Legacy Name", 2), ("ISS-1", "Ali Khan", 11),
    ]


def _report(deadline, submitted=None, kind="initial_notification"):
    from app.models.enums import RegulatoryReportStatus, RegulatoryReportType

    return SimpleNamespace(report_type=RegulatoryReportType(kind), deadline=deadline, submitted_at=submitted,
                           reference="", status=RegulatoryReportStatus.submitted if submitted else RegulatoryReportStatus.pending)


def _incident(ref, sev="high", reportable=False, occurred=None, detected=None, resolved=None, near_miss=False):
    return SimpleNamespace(id=uuid.uuid4(), reference=ref, title=f"Incident {ref}", severity=sev,
                           is_reportable=reportable, regulator="SBP", near_miss=near_miss, occurred_at=occurred,
                           detected_at=detected, contained_at=None, resolved_at=resolved, created_at=detected)


def test_incident_summary_rates_only_decided_notifications():
    now = at(2026, 9, 12, 12)
    on_time = _incident("INC-1", "critical", True, at(2026, 8, 1, 0), at(2026, 8, 1, 2), at(2026, 8, 2, 2))
    late = _incident("INC-2", "high", True, at(2026, 8, 5), at(2026, 8, 5, 4))
    pending = _incident("INC-3", "high", True, detected=at(2026, 9, 12, 6))
    plain = _incident("INC-4", "low", occurred=at(2026, 8, 9), detected=at(2026, 8, 9, 6), near_miss=True)
    reports = {
        on_time.id: [_report(at(2026, 8, 2, 2), submitted=at(2026, 8, 1, 20))],
        late.id: [_report(at(2026, 8, 6, 4), submitted=at(2026, 8, 7))],
        pending.id: [_report(at(2026, 9, 13, 6))],
    }
    s = bp.incident_summary([on_time, late, pending, plain], reports, now)
    assert (s["total"], s["reportable"], s["near_misses"]) == (4, 3, 1)
    assert (s["on_time"], s["late"], s["pending"]) == (1, 1, 1)
    assert s["on_time_rate"] == 50.0
    assert s["by_severity"] == {"critical": 1, "high": 2, "medium": 0, "low": 1}
    # MTTD over the incidents with both times (2 h, 4 h, 6 h); MTTR over the one resolved.
    assert (s["mttd_hours"], s["mttd_n"]) == (4.0, 3)
    assert (s["mttr_hours"], s["mttr_n"]) == (24.0, 1)
    assert [r["reference"] for r in s["rows"]] == ["INC-1", "INC-2", "INC-3"]
    assert [r["on_time"] for r in s["rows"]] == [True, False, None]


def test_no_decided_notification_means_no_rate():
    s = bp.incident_summary([_incident("INC-1")], {}, at(2026, 9, 12))
    assert s["on_time_rate"] is None and s["reportable"] == 0


def _kri(ref, status, owner_id=None):
    return SimpleNamespace(reference=ref, name=f"KRI {ref}", status=KriStatus(status), current_value=5,
                           unit="%", warning_threshold=3, limit_threshold=4, lower_bound=None, upper_bound=None,
                           direction="higher_is_worse", owner="text owner", owner_id=owner_id,
                           last_measured_date=date(2026, 9, 1))


def test_kri_summary_lists_red_then_amber():
    uid = uuid.uuid4()
    s = bp.kri_summary([_kri("K-3", "amber"), _kri("K-1", "green"), _kri("K-2", "red", uid), _kri("K-4", "no_data")],
                       {uid: "Sara"})
    assert (s["red"], s["amber"], s["green"], s["no_data"]) == (1, 1, 1, 1)
    assert [(r["reference"], r["owner"]) for r in s["rows"]] == [("K-2", "Sara"), ("K-3", "text owner")]


def test_third_parties_critical_and_high_with_lapsing_certificates():
    today = date(2026, 9, 12)
    acme, beta, gamma = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    vendors = [
        SimpleNamespace(id=acme, name="Acme", criticality="critical", inherent_tier="high", status="active",
                        next_review_date=date(2026, 8, 1)),
        SimpleNamespace(id=beta, name="Beta", criticality="high", inherent_tier=None, status="active",
                        next_review_date=None),
        SimpleNamespace(id=gamma, name="Gamma", criticality="medium", inherent_tier="medium", status="active",
                        next_review_date=None),
    ]
    certs = [
        SimpleNamespace(vendor_id=acme, vendor_name="Acme", cert_type="ISO 27001", expires_on=date(2026, 9, 1)),
        SimpleNamespace(vendor_id=beta, vendor_name="Beta", cert_type="SOC 2 Type II", expires_on=date(2026, 10, 30)),
        SimpleNamespace(vendor_id=gamma, vendor_name="Gamma", cert_type="PCI DSS", expires_on=date(2027, 6, 1)),
    ]
    s = bp.third_party_summary(vendors, certs, today)
    assert (s["critical"], s["high"], s["expired_certs"], s["expiring_certs"]) == (1, 1, 1, 1)
    assert [r["name"] for r in s["rows"]] == ["Acme", "Beta"]
    assert s["rows"][0]["review_overdue"] and s["rows"][0]["certs_lapsing"] == 1
    assert [(c["vendor"], c["state"]) for c in s["certifications"]] == [("Acme", "expired"), ("Beta", "expiring")]


def test_failed_tests_count_reviewed_ones_only():
    t = lambda result, review, ref="A.8.5": SimpleNamespace(  # noqa: E731
        result=result, review_status=review, conducted_date=date(2026, 8, 1), control_reference=ref,
        control_name="Secure authentication", is_key=True, issue_reference="ISS-9")
    s = bp.failed_tests_summary([
        t("failed", "reviewed"), t("failed", "legacy", "A.8.2"), t("failed", "pending"), t("failed", "returned"),
        t("passed_with_exceptions", "reviewed"), t("passed_with_exceptions", "pending"),
    ])
    assert len(s["failed"]) == 2
    assert s["failed_pending_review"] == 2
    assert s["exceptions"] == 1
    assert s["failed"][0]["control"].startswith("A.8") and s["failed"][0]["issue"] == "ISS-9"


# ------------------------------------------------------------- overview sections ---
def _overview():
    comp = lambda key, value, pop: SimpleNamespace(label=key.title(), value=value, weight=0.25, detail=f"{key} detail")  # noqa: E731,ARG005
    return SimpleNamespace(
        health=SimpleNamespace(score=64, band="elevated", components=[comp("tolerance", 90.0, 10), comp("discipline", None, 0)],
                               coverage=SimpleNamespace(scored=1, total=2, weight_pct=50.0)),
        assurance=SimpleNamespace(total=12, effective=4, partially_effective=2, ineffective=1, not_assessed=3,
                                  not_operating=2, tests_overdue=1, last_test_failed=1),
        posture=SimpleNamespace(total_risks=10, breach=2, elevated=3, within_appetite=5, appetite_score=6,
                                tolerance_score=12, by_category=[], top_risks=[]),
        compliance=SimpleNamespace(overall_assured_pct=40.0, frameworks=[SimpleNamespace(
            name="ISO/IEC 27001:2022", applicable=100, assured=40, unassessed=50, failing=2, unmapped=8, gaps=60,
            compliant_pct=12.5)]),
        incidents=SimpleNamespace(open=3, reportable_open=1),
        kris=SimpleNamespace(red=2),
        actions=[SimpleNamespace(label="2 risks above tolerance", count=2)],
    )


def test_summary_uses_the_dashboard_score_and_its_coverage():
    s = bp.summary_from_overview(_overview())
    assert (s["score"], s["band"], s["scored"], s["total"], s["weight_pct"]) == (64, "elevated", 1, 2, 50.0)
    assert ("Risks above tolerance", "2 of 10") in s["headlines"]
    assert ("Operating controls effective or partially effective", "6 of 10") in s["headlines"]
    assert s["actions"] == [{"label": "2 risks above tolerance", "count": 2}]


def test_an_organisation_with_nothing_to_score_has_no_score_in_the_pack():
    o = _overview()
    o.health.band = "no_data"
    assert bp.summary_from_overview(o)["score"] is None


def test_appetite_falls_back_to_the_organisation_row():
    a = bp.appetite_from_overview(_overview())
    assert a["rows"] == [{"label": "All categories (organisation appetite)", "appetite": 6, "tolerance": 12,
                          "risks": 10, "within": 5, "elevated": 3, "breach": 2}]


def test_compliance_rows_carry_the_assured_share():
    c = bp.compliance_from_overview(_overview())
    assert c["rows"][0]["assured_pct"] == 40.0 and c["overall_assured_pct"] == 40.0


async def test_the_dashboard_overview_is_called_with_its_own_defaults(monkeypatch):
    """The pack calls the dashboard endpoint itself; a query parameter added to it later
    receives its declared default instead of a FastAPI ``Query`` object."""
    from fastapi import Query

    from app.api.v1 import dashboard

    seen = {}

    async def fake(db, user, days: int = Query(default=30, ge=7), level: int | None = Query(default=None), flag: bool = True):
        seen.update(db=db, tenant=user.tenant_id, days=days, level=level, flag=flag)
        return "overview"

    monkeypatch.setattr(dashboard, "get_overview", fake)
    tid = uuid.uuid4()
    assert await bp.dashboard_overview("DB", tid, 73) == "overview"
    assert seen == {"db": "DB", "tenant": tid, "days": 73, "level": None, "flag": True}
    # A person generating the pack is the dashboard's caller; the scheduler is a system
    # actor holding every permission, with no user id.
    person = SimpleNamespace(tenant_id=tid)
    captured = {}

    async def who(db, user, days: int = Query(default=30)):
        captured["user"] = user

    monkeypatch.setattr(dashboard, "get_overview", who)
    await bp.dashboard_overview("DB", tid, 30, person)
    assert captured["user"] is person
    await bp.dashboard_overview("DB", tid, 30)
    system = captured["user"]
    assert system.id is None and "risk:read" in system.permission_codes and system.tenant_id == tid


# ------------------------------------------------------------------ views & files ---
def _pack(**sections):
    return {
        "title": "Board pack: Risk & Audit <Committee>",
        "sections": list(sections),
        "cover": {"organisation": "Demo Bank", "committee": "Risk & Audit", "committee_reference": "CMT-001",
                  "meeting": "Q3", "meeting_reference": "MTG-002", "meeting_date": date(2026, 9, 24),
                  "period_start": date(2026, 7, 1), "period_end": date(2026, 9, 12), "as_of": date(2026, 9, 12),
                  "generated_at": at(2026, 9, 12, 9), "generated_by": "Jane", "timezone": "Asia/Karachi",
                  "tz": KHI, "date_format": "DD/MM/YYYY"},
        **sections,
    }


MOVEMENT = {"new": [{"reference": f"R-{i}", "title": "New risk", "score": 9, "when": at(2026, 8, 1)} for i in range(70)],
            "closed": [], "rescored": [], "counts": {"new": 70, "closed": 0, "rescored": 0, "up": 0, "down": 0}}


def test_the_pdf_lists_up_to_the_limit_and_the_spreadsheet_everything():
    pack = _pack(movement=MOVEMENT)
    pdf_view = bp.section_views(pack)[0]
    xlsx_view = bp.section_views(pack, limit=None)[0]
    assert len(pdf_view.tables[0].rows) == bp.LIST_LIMIT and pdf_view.tables[0].total == 70
    assert len(xlsx_view.tables[0].rows) == 70
    assert ("Added", "70") in pdf_view.kpis


def test_cover_says_what_the_pack_is():
    facts = dict(bp.cover_facts(_pack(movement=MOVEMENT)))
    assert facts["Committee"] == "CMT-001 Risk & Audit"
    assert facts["Meeting"] == "MTG-002 Q3, 24/09/2026"
    assert facts["Period"] == "01/07/2026 to 12/09/2026"
    assert facts["Generated"] == "12/09/2026 14:00 (Asia/Karachi)"
    assert facts["Sections"] == "Risk movement"


def test_the_spreadsheet_has_a_cover_and_one_sheet_per_section():
    from openpyxl import load_workbook

    pack = _pack(movement=MOVEMENT, kris={"green": 0, "amber": 0, "red": 0, "no_data": 0, "rows": []})
    wb = load_workbook(io.BytesIO(bp.to_xlsx(pack)))
    assert wb.sheetnames == ["Cover", "Risk movement", "Key risk indicators"]
    assert wb["Cover"]["A1"].value == pack["title"]
    values = [c.value for row in wb["Risk movement"].iter_rows() for c in row]
    assert values.count("New risk") == 70
    assert "No KRI is at red or amber." in [c.value for row in wb["Key risk indicators"].iter_rows() for c in row]


def test_the_pdf_renders_and_escapes_markup():
    pdf = bp.to_pdf(_pack(movement=MOVEMENT))
    assert pdf.startswith(b"%PDF")


def test_every_section_renders():
    """A pack with every section, each with rows, renders in both formats."""
    pack = _pack(
        summary=bp.summary_from_overview(_overview()),
        appetite=bp.appetite_from_overview(_overview()),
        top_risks={"rows": [{"reference": "R-1", "title": "Outage", "score": 20, "severity": "critical",
                             "appetite_status": "breach", "owner": "", "status": "assessed", "trend_text": "Up from 12"}]},
        movement=MOVEMENT,
        assurance={"total": 1, "effective": 1, "partially_effective": 0, "ineffective": 0, "not_assessed": 0,
                   "not_operating": 0, "tests_overdue": 0, "last_test_failed": 0,
                   "failed": [{"control": "A.8.5", "conducted": date(2026, 8, 1), "key": True, "issue": ""}],
                   "failed_pending_review": 0, "exceptions": 0},
        compliance=bp.compliance_from_overview(_overview()),
        issues=bp.issue_summary([_issue("ISS-1", "high", date(2026, 9, 1))], {}, date(2026, 9, 12)),
        incidents=bp.incident_summary([_incident("INC-1", reportable=True, detected=at(2026, 8, 1))],
                                      {}, at(2026, 9, 12)),
        kris=bp.kri_summary([_kri("K-1", "red")]),
        third_parties=bp.third_party_summary([], [], date(2026, 9, 12)),
    )
    assert [v.key for v in bp.section_views(pack)] == list(bp.SECTION_KEYS)
    assert bp.to_pdf(pack).startswith(b"%PDF")
    assert len(bp.to_xlsx(pack)) > 1000


def test_sheet_names_are_excel_safe_and_unique():
    used: set = set()
    assert bp.sheet_name("Risks: in/out [Q3]?", used) == "Risks inout Q3"
    assert bp.sheet_name("Risks: in/out [Q3]?", used) == "Risks inout Q3 (2)"
    assert len(bp.sheet_name("x" * 40, used)) == 31


def test_files_are_filed_on_the_meeting_then_the_committee_then_the_pack():
    pack_id = uuid.uuid4()
    committee, meeting = SimpleNamespace(id=uuid.uuid4()), SimpleNamespace(id=uuid.uuid4())
    assert bp.attach_target(pack_id, committee, meeting) == ("committee_meeting", meeting.id)
    assert bp.attach_target(pack_id, committee, None) == ("committee", committee.id)
    assert bp.attach_target(pack_id) == ("board_pack", pack_id)
    assert bp.file_stem("Board pack: BRMC — Q3!", date(2026, 9, 12)) == "board-pack-brmc-q3-2026-09-12"


# ------------------------------------------------------------- generate (fake db) ---
class _Nested:
    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False  # let the exception through, as a savepoint does after rolling back


class FakeDB:
    def __init__(self):
        self.added = []

    def begin_nested(self):
        return _Nested()

    def add(self, obj):
        self.added.append(obj)

    async def flush(self):
        pass


def _org():
    return bp.OrgContext(tenant_id=uuid.uuid4(), name="Demo Bank", tz=KHI, timezone="Asia/Karachi",
                         date_format="DD/MM/YYYY", fiscal_start_month=1, today=date(2026, 9, 12),
                         now=at(2026, 9, 12, 9))


async def test_generate_files_both_formats_and_audits_the_person(monkeypatch):
    from app.services import audit

    stored, audited = [], []

    async def fake_build(db, org, **kw):
        return {"title": kw["title"], "sections": kw["sections"], "cover": {}}

    async def fake_store(db, tenant_id, **kw):
        f = SimpleNamespace(id=uuid.uuid4(), storage_key=f"k/{kw['filename']}", filename=kw["filename"], **{
            "entity_type": kw["entity_type"], "entity_id": kw["entity_id"]})
        stored.append(f)
        return f

    async def fake_record(db, **kw):
        audited.append(kw)

    monkeypatch.setattr(bp, "build_pack", fake_build)
    monkeypatch.setattr(bp, "to_pdf", lambda data: b"%PDF-1.4")
    monkeypatch.setattr(bp, "to_xlsx", lambda data: b"PK")
    monkeypatch.setattr(bp, "store_file", fake_store)
    monkeypatch.setattr(audit, "record", fake_record)
    user = SimpleNamespace(id=uuid.uuid4(), email="jane@bank.pk", full_name="Jane", tenant_id=uuid.uuid4())
    committee = SimpleNamespace(id=uuid.uuid4(), name="BRMC", reference="CMT-001")
    meeting = SimpleNamespace(id=uuid.uuid4(), title="Q3 sitting", reference="MTG-001", meeting_date=date(2026, 9, 24))
    db = FakeDB()
    row = await bp.generate(db, _org(), actor=user, committee=committee, meeting=meeting, sections=["kris"])
    assert row.status == bp.READY and row.error == ""
    assert row.pdf_file_id == stored[0].id and row.xlsx_file_id == stored[1].id
    assert {f.entity_type for f in stored} == {"committee_meeting"} and stored[0].entity_id == meeting.id
    assert (row.period_start, row.period_end, row.sections) == (date(2026, 7, 1), date(2026, 9, 12), ["kris"])
    assert row.title == "Board pack: BRMC — Q3 sitting" and row.generated_by_id == user.id
    assert row in db.added and row.generated_at is not None
    assert audited[0]["actor"] is user and audited[0]["entity_type"] == "committee"
    assert audited[0]["action"] == "board_pack" and audited[0]["changes"]["status"] == "ready"


async def test_a_failed_generation_is_kept_without_files(monkeypatch):
    from app.services import audit, storage

    deleted, audited = [], []
    calls = {"n": 0}

    async def fake_build(db, org, **kw):
        return {}

    async def fake_store(db, tenant_id, **kw):
        calls["n"] += 1
        if calls["n"] == 2:
            raise RuntimeError("disk full")
        return SimpleNamespace(id=uuid.uuid4(), storage_key="k/1", filename=kw["filename"])

    async def fake_system(db, **kw):
        audited.append(kw)

    monkeypatch.setattr(bp, "build_pack", fake_build)
    monkeypatch.setattr(bp, "to_pdf", lambda data: b"%PDF")
    monkeypatch.setattr(bp, "to_xlsx", lambda data: b"PK")
    monkeypatch.setattr(bp, "store_file", fake_store)
    monkeypatch.setattr(storage, "delete_object", deleted.append)
    monkeypatch.setattr(audit, "record_system", fake_system)
    db = FakeDB()
    row = await bp.generate(db, _org(), actor=None, reason="7 day(s) before the meeting")
    assert row.status == bp.FAILED and row.error == "disk full"
    assert row.pdf_file_id is None and row.xlsx_file_id is None
    assert deleted == ["k/1"]  # the PDF already written is removed
    assert audited[0]["entity_type"] == "board_pack" and "could not be generated" in audited[0]["summary"]
    assert row.generated_by_id is None


async def test_a_bad_request_writes_nothing():
    db = FakeDB()
    with pytest.raises(bp.PackError):
        await bp.generate(db, _org(), period_end=date(2026, 9, 13))
    with pytest.raises(bp.PackError):
        await bp.generate(db, _org(), sections=["nope"])
    assert db.added == []


def test_scheduled_packs_name_the_scheduler():
    assert bp.actor_label(None) == bp.SCHEDULER_NAME
    assert bp.actor_label(SimpleNamespace(full_name="", email="a@b.pk")) == "a@b.pk"


def test_error_text_prefers_an_http_detail():
    from fastapi import HTTPException

    assert bp.error_text(HTTPException(status_code=501, detail="PDF export requires reportlab")) == "PDF export requires reportlab"
    assert bp.error_text(ValueError("x" * 2000)) == "x" * 1000


# ------------------------------------------------------------- committee setting ---
def test_the_auto_generate_setting_is_bounded():
    from pydantic import ValidationError

    from app.schemas.governance import CommitteeCreate, CommitteeUpdate

    assert CommitteeCreate(name="BRMC", board_pack_days_before=7).board_pack_days_before == 7
    assert CommitteeCreate(name="BRMC").board_pack_days_before is None
    assert CommitteeUpdate(board_pack_days_before=None).model_dump(exclude_unset=True) == {"board_pack_days_before": None}
    for bad in (0, 91, -1):
        with pytest.raises(ValidationError):
            CommitteeCreate(name="BRMC", board_pack_days_before=bad)


async def test_meetings_are_loaded_through_their_committee():
    """``Meeting`` has no archive flag; the loader used to filter on one, so opening,
    editing or adding a decision to any meeting failed. It now hides the meetings of an
    archived committee instead."""
    from sqlalchemy.dialects import postgresql

    from app.api.v1 import governance
    from app.models.governance import Meeting

    seen = {}

    class DB:
        async def scalar(self, stmt):
            seen["sql"] = str(stmt.compile(dialect=postgresql.dialect()))
            return "meeting"

    assert not hasattr(Meeting, "deleted")
    assert await governance._load_meeting(DB(), uuid.uuid4()) == "meeting"
    assert "JOIN committees" in seen["sql"] and "committees.deleted" in seen["sql"]


# ------------------------------------------------- loaders build valid PostgreSQL ---
class CompilingDB:
    """Compiles every statement for PostgreSQL (so a wrong column or construct fails
    here, not on a bank's server) and answers with no rows."""

    def __init__(self):
        self.sql: list[str] = []

    def _compile(self, stmt):
        from sqlalchemy.dialects import postgresql

        self.sql.append(str(stmt.compile(dialect=postgresql.dialect())))

    async def execute(self, stmt, *a, **kw):
        self._compile(stmt)
        return SimpleNamespace(all=lambda: [])

    async def scalars(self, stmt, *a, **kw):
        self._compile(stmt)
        return SimpleNamespace(all=lambda: [])

    async def scalar(self, stmt, *a, **kw):
        self._compile(stmt)
        return None


async def test_every_loader_builds_valid_sql_and_an_empty_pack_renders(monkeypatch):
    async def fake_overview(db, tenant_id, days, viewer=None):
        return _overview()

    monkeypatch.setattr(bp, "dashboard_overview", fake_overview)
    db = CompilingDB()
    org = _org()
    pack = await bp.build_pack(db, org, period_start=date(2026, 7, 1), period_end=date(2026, 9, 12),
                               sections=list(bp.SECTION_KEYS), title="Board pack")
    assert all(k in pack for k in bp.SECTION_KEYS)
    joined = "\n".join(db.sql)
    assert "audit_logs.changes ? " in joined          # re-assessed per the trail
    assert "issue_due_date_changes" in joined and "regulatory_reports" not in joined  # no reportable incidents
    assert pack["movement"]["counts"] == {"new": 0, "closed": 0, "rescored": 0, "up": 0, "down": 0}
    assert pack["incidents"]["on_time_rate"] is None
    assert bp.to_pdf(pack).startswith(b"%PDF") and bp.to_xlsx(pack)


async def test_history_queries_compile():
    db = CompilingDB()
    snaps, assessed = await bp._history_before(db, [uuid.uuid4()], START)
    assert snaps == {} and assessed == {}
    assert any("DISTINCT ON (record_versions.entity_id)" in q for q in db.sql)
    assert any("DISTINCT ON (audit_logs.entity_id)" in q for q in db.sql)


# ------------------------------------------------------------ the scheduler step ---
class SweepDB(CompilingDB):
    def __init__(self, *, committees, meetings, existing=None, lock=True, enabled_modules=None):
        super().__init__()
        self.committees, self.meetings = committees, meetings
        self.existing, self.lock, self.enabled_modules = existing, lock, enabled_modules

    async def execute(self, stmt, *a, **kw):
        self._compile(stmt)
        sql = self.sql[-1]
        rows = self.committees if "FROM committees" in sql else self.meetings if "FROM committee_meetings" in sql else []
        return SimpleNamespace(all=lambda: rows)

    async def scalar(self, stmt, *a, **kw):
        sql = str(stmt) if not hasattr(stmt, "compile") else None
        if sql is None:
            self._compile(stmt)
            sql = self.sql[-1]
        if "pg_try_advisory_xact_lock" in sql:
            return self.lock
        if "FROM board_packs" in sql:
            return self.existing
        if "enabled_modules" in sql:
            return self.enabled_modules
        return None


@pytest.fixture
def sweep(monkeypatch):
    from app.services import modules

    made = []

    async def fake_org(db, tenant_id):
        return _org()

    async def fake_generate(db, org, **kw):
        made.append(kw)
        return SimpleNamespace(**kw)

    monkeypatch.setattr(bp, "org_context", fake_org)
    monkeypatch.setattr(bp, "generate", fake_generate)
    monkeypatch.setattr(modules, "is_enabled", lambda key: True)
    return made


COMMITTEE = SimpleNamespace(id=uuid.uuid4(), name="BRMC", reference="CMT-001", board_pack_days_before=7)


def _meeting(day):
    return SimpleNamespace(id=uuid.uuid4(), title="Q3 sitting", reference="MTG-001", meeting_date=day)


async def test_the_scheduler_generates_a_pack_for_a_meeting_now_in_its_window(sweep):
    m = _meeting(date(2026, 9, 18))
    out = await bp.generate_due_packs(SweepDB(committees=[COMMITTEE], meetings=[m]), uuid.uuid4())
    assert len(out) == 1 and sweep[0]["meeting"].id == m.id and sweep[0]["committee"].name == "BRMC"
    assert sweep[0]["reason"] == "7 day(s) before the meeting" and "actor" not in sweep[0]


async def test_the_scheduler_skips_what_is_not_due_or_already_done(sweep):
    far = _meeting(date(2026, 10, 30))       # outside the 7-day window
    soon = _meeting(date(2026, 9, 15))
    assert await bp.generate_due_packs(SweepDB(committees=[COMMITTEE], meetings=[far]), uuid.uuid4()) == []
    assert await bp.generate_due_packs(
        SweepDB(committees=[COMMITTEE], meetings=[soon], existing=uuid.uuid4()), uuid.uuid4()) == []  # pack in window
    assert await bp.generate_due_packs(
        SweepDB(committees=[COMMITTEE], meetings=[soon], lock=False), uuid.uuid4()) == []  # another worker has it
    assert await bp.generate_due_packs(
        SweepDB(committees=[COMMITTEE], meetings=[soon], enabled_modules=["risk"]), uuid.uuid4()) == []  # module off
    assert sweep == []

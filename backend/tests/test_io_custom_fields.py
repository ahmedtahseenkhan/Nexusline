"""Custom fields in the import template, the export and the import.

An organisation's custom fields are part of its register: the template must offer them,
the export must carry their values, and a filled template or edited export must import
back onto the same fields with the same validation the record form applies.
"""
import io
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook

from app.api.v1.dataio import _custom_cell, _custom_columns, _typed_custom
from app.models.custom_field import CustomField
from app.models.enums import CustomFieldType
from app.services import csv_io, xlsx_io
from app.services.import_mapping import load_xlsx
from app.services.import_registry import REGISTRY

RISKS = REGISTRY["risks"]


class _FakeDb:
    """Just enough of an AsyncSession for ``_custom_columns``: returns the given rows."""

    def __init__(self, rows):
        self.rows = rows
        self.statement = None

    async def scalars(self, statement):
        self.statement = statement
        return self

    def all(self):
        return self.rows


def _field(label, field_type=CustomFieldType.text, *, options="", required=False, model="risk"):
    return CustomField(
        id=uuid.uuid4(), model=model, label=label, field_type=field_type, options=options,
        required=required, help_text="", order_index=0, enabled=True,
    )


@pytest.mark.asyncio
async def test_custom_fields_become_typed_columns_after_the_built_in_ones():
    fields = [
        _field("Cost centre"),
        _field("Board approved", CustomFieldType.checkbox),
        _field("Tier", CustomFieldType.select, options="Tier 1\nTier 2\n", required=True),
        _field("Go-live", CustomFieldType.date),
    ]
    cols = [col for col, _ in await _custom_columns(_FakeDb(fields), RISKS)]
    assert [c.header for c in cols] == ["Cost centre", "Board approved", "Tier", "Go-live"]
    assert [c.kind for c in cols] == ["text", "bool", "enum", "date"]
    assert cols[2].enum_values == ["Tier 1", "Tier 2"] and cols[2].required


@pytest.mark.asyncio
async def test_a_label_that_repeats_a_built_in_heading_is_suffixed():
    builtin = RISKS.columns[0].header
    fields = [_field(builtin.upper()), _field("Owner note"), _field("owner  NOTE")]
    headers = [col.header for col, _ in await _custom_columns(_FakeDb(fields), RISKS)]
    assert headers == [f"{builtin.upper()} (custom)", "Owner note", "owner  NOTE (custom)"]


@pytest.mark.asyncio
async def test_a_register_without_custom_fields_adds_no_columns():
    db = _FakeDb([_field("x")])
    assert await _custom_columns(db, REGISTRY["fx-rates"]) == []
    assert db.statement is None  # no query at all


@pytest.mark.asyncio
async def test_it_and_information_asset_exports_query_their_own_fields():
    for resource, key in (("it-assets", "it_asset"), ("information-assets", "information_asset")):
        db = _FakeDb([])
        await _custom_columns(db, REGISTRY[resource])
        compiled = db.statement.compile(compile_kwargs={"literal_binds": True})
        assert f"'{key}'" in str(compiled)


# ------------------------------------------------------------ cell values ---
def _col(kind, options=None):
    from app.services.import_registry import Column

    return Column(header="Field", field="cf:x", kind=kind, enum_values=options)


def test_import_cells_are_validated_and_normalised():
    assert _custom_cell(_col("bool"), "Yes") == "true"
    assert _custom_cell(_col("bool"), "0") == "false"
    assert _custom_cell(_col("date"), "2026-03-01") == "2026-03-01"
    assert _custom_cell(_col("enum", ["Tier 1", "Tier 2"]), "tier 2") == "Tier 2"
    assert _custom_cell(_col("float"), " 12.5 ") == "12.5"
    assert _custom_cell(_col("text"), "   ") == ""


@pytest.mark.parametrize(
    ("kind", "options", "raw"),
    [("enum", ["A", "B"], "C"), ("date", None, "01/03/2026"), ("float", None, "lots"), ("bool", None, "maybe")],
)
def test_a_bad_cell_names_its_column(kind, options, raw):
    with pytest.raises(ValueError, match="^Field: "):
        _custom_cell(_col(kind, options), raw)


def test_export_types_stored_values_for_excel_but_keeps_legacy_text():
    assert _typed_custom(_col("date"), "2026-03-01") == date(2026, 3, 1)
    assert _typed_custom(_col("bool"), "true") is True
    assert _typed_custom(_col("float"), "3") == 3.0
    assert _typed_custom(_col("date"), "not a date") == "not a date"


# ------------------------------------------------------------- workbook ---
@pytest.mark.asyncio
async def test_excel_template_round_trips_through_the_importer():
    fields = [_field("Tier", CustomFieldType.select, options="Tier 1\nTier 2", required=True)]
    custom = await _custom_columns(_FakeDb(fields), RISKS)
    columns = [*RISKS.columns, *(c for c, _ in custom)]
    data = xlsx_io.build_workbook(columns, [], title="Risks", examples=csv_io.example_row(columns))

    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Risks", "Lists", "Guide"]
    assert wb["Lists"].sheet_state == "hidden"
    sheet = wb["Risks"]
    headers = [c.value for c in sheet[1]]
    assert headers == [c.header for c in columns]
    tier_letter = sheet.cell(row=1, column=len(columns)).column_letter
    assert any(
        f"{tier_letter}2" in str(dv.sqref) for dv in sheet.data_validations.dataValidation
    ), "the select custom field gets an in-cell dropdown"
    guide = [[c.value for c in row] for row in wb["Guide"].iter_rows(min_row=2)]
    assert ["Tier", "Yes", "One of the listed values", "Tier 1, Tier 2", "Tier 1", "Custom field"] in guide

    # Fill one row as a user would and read it back the way the importer does.
    sheet.cell(row=2, column=1, value="Imported risk")
    sheet.cell(row=2, column=len(columns), value="Tier 2")
    buf = io.BytesIO()
    wb.save(buf)
    table = load_xlsx(buf.getvalue())
    assert table.sheet == "Risks"
    assert table.headers == [c.header for c in columns]
    assert table.rows[0][-1] == "Tier 2"


def test_excel_export_writes_real_dates_and_booleans():
    from app.services.import_registry import Column

    cols = [Column(header="Due", field="due", kind="date"), Column(header="Flag", field="f", kind="bool")]
    data = xlsx_io.build_workbook(cols, [{"Due": date(2026, 1, 2), "Flag": True}], title="X")
    sheet = load_workbook(io.BytesIO(data))["X"]
    assert sheet["A2"].is_date and sheet["B2"].value is True
    assert load_xlsx(data).rows[0] == ["2026-01-02", "true"]


def test_exports_write_enum_values_not_python_names():
    """A str-Enum renders as "Criticality.medium" through ``str()``; an export carrying
    that would never import back."""
    from app.models.enums import Criticality
    from app.services.import_registry import Column

    assert csv_io.export_csv([{"c": Criticality.medium}], ["c"]) == "c\nmedium\n"
    data = xlsx_io.build_workbook(
        [Column(header="c", field="c", kind="enum", enum_values=["medium"])],
        [{"c": Criticality.medium}], title="X",
    )
    assert load_xlsx(data).rows[0] == ["medium"]


def test_timezone_aware_datetimes_export_as_utc():
    """Excel rejects tz-aware datetimes outright (an incident's detected_at 500'd)."""
    from datetime import datetime, timedelta, timezone

    from app.services.import_registry import Column

    karachi = timezone(timedelta(hours=5))
    data = xlsx_io.build_workbook(
        [Column(header="At", field="at")],
        [{"At": datetime(2026, 1, 2, 15, 30, tzinfo=karachi)}], title="X",
    )
    assert load_xlsx(data).rows[0] == ["2026-01-02T10:30:00"]

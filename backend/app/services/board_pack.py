"""Board packs — the committee's view of risk, assurance and compliance for a period,
kept as the PDF and spreadsheet the committee actually saw.

A pack is built from the same numbers the dashboard shows, so the two can never
disagree: the governance-health score, appetite by category, top risks, control
assurance and compliance come from the dashboard overview itself (which reads
``governance_health``, ``risk_settings.load_appetite_book`` and ``control_assurance``),
and the period sections below add what a committee asks about the period:

* **Risk movement** — risks added, closed and re-scored between the period's dates.
* **Top-risk trend** — each top risk's score now against its score when the period
  began. A risk's scores are stamped (``last_assessed_at``) on every change, so a risk
  not re-scored since the period began is *unchanged* for certain; for one re-scored in
  the period the score at the start comes from the record's version history (and, where
  a residual was re-assessed after the last snapshot, from that assessment's audit
  entry). A risk with no history before the period is shown as such, never guessed.
* **Failed control tests** conducted in the period, counting reviewed tests only (an
  unreviewed test changes nothing — Phase 2) and saying how many await review.
* **Open issues** by severity, overdue and how many had their due date moved.
* **Incidents** detected in the period: reportable ones, the regulator-notification
  on-time rate, and mean time to detect and to resolve.
* **KRIs** at red and amber, and **third parties** rated critical or high with their
  lapsing certifications.

Position figures are *as at* the day the pack is generated; movement and incidents
cover the period. The pack says so on its cover.

Everything that decides a number is pure (the period, sections, trend, the section
views both renderers print) and unit-tested; the async functions at the bottom load
the rows for one organisation, store the files and write the ``BoardPack`` row.

**Phase 4B.**

* **Sign-off.** A pack is generated as a *draft*; someone other than the people who
  shaped it (generated it or wrote its commentary) marks it *reviewed*, and it is then
  *released* — to the committee's member users, each notified in the app and by e-mail
  with a link. Four-eyes follows the ``board_pack`` / ``release`` dual-control key.
  Changing the commentary of a reviewed pack returns it to draft; a released pack is
  frozen (generate a new one to change it).
* **Commentary.** A narrative per section (the CRO's words), printed at the top of the
  section in the PDF and the spreadsheet. The figures a pack was built from are kept on
  the row (``content``), so adding commentary re-renders the same numbers, not today's.
* **Saved sections.** A committee keeps its own section choice and order
  (``Committee.board_pack_sections``); packs generated for it — by hand without a choice,
  or by the scheduler — use it.
* **Charts.** The appetite section carries the heat map of current exposure and the
  appetite position over the last four quarter ends; the KRI section each listed KRI's
  recent readings.
* **Branding.** Logo, primary colour, cover title and the classification printed on
  every page (``BoardPackBranding``; "Confidential" unless set).
* **Past periods.** A pack whose period ended before today takes its position figures
  from the period snapshot nearest the period end (``services/snapshots.py``) and says so
  on the cover and in each section; with no snapshot near enough they are as at
  generation, and the cover says that instead.
"""
from __future__ import annotations

import io
import logging
import uuid
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from typing import Any, Iterable, Mapping, Sequence
from xml.sax.saxutils import escape

logger = logging.getLogger("nexusline.board_pack")

# ---------------------------------------------------------------------------
# Sections
# ---------------------------------------------------------------------------
SECTIONS: tuple[tuple[str, str], ...] = (
    ("summary", "Executive summary"),
    ("appetite", "Appetite by category"),
    ("top_risks", "Top risks"),
    ("movement", "Risk movement"),
    ("assurance", "Control assurance"),
    ("compliance", "Compliance by framework"),
    ("issues", "Open issues"),
    ("incidents", "Incidents"),
    ("kris", "Key risk indicators"),
    ("third_parties", "Third parties"),
)
SECTION_KEYS: tuple[str, ...] = tuple(k for k, _ in SECTIONS)
SECTION_TITLES: dict[str, str] = dict(SECTIONS)
#: Sections whose figures come from the dashboard overview (the KRI section reads the
#: same ``KeyRiskIndicator.status`` rule the dashboard does, with its own list).
OVERVIEW_SECTIONS = frozenset({"summary", "appetite", "top_risks", "assurance", "compliance"})

#: BoardPack.status values.
READY, FAILED = "ready", "failed"
#: BoardPack.review_state values (phase 4B).
DRAFT, REVIEWED, RELEASED = "draft", "reviewed", "released"
REVIEW_STATES = (DRAFT, REVIEWED, RELEASED)
#: The dual-control key that governs reviewing and releasing a pack.
DUAL_CONTROL_MODULE, DUAL_CONTROL_ACTION = "board_pack", "release"
#: Longest commentary per section, in characters.
COMMENTARY_MAX = 6000
#: Quarter ends the appetite trend chart covers.
TREND_QUARTERS = 4
#: KRI readings drawn per KRI.
KRI_TREND_READINGS = 8
#: KRIs given a trend chart in the PDF.
KRI_TREND_CHARTS = 6

#: Rows listed per table in a pack (the counts above each table are never capped).
LIST_LIMIT = 50
#: Actor named on a pack the scheduler generated.
SCHEDULER_NAME = "Scheduler"

#: Issue states that are closed (the issue module's own definition).
CLOSED_ISSUE_STATES = ("closed", "remediated", "risk_accepted")
#: Severity order, worst first.
SEVERITIES = ("critical", "high", "medium", "low")

DATE_FORMATS = {
    "DD/MM/YYYY": "%d/%m/%Y",
    "MM/DD/YYYY": "%m/%d/%Y",
    "YYYY-MM-DD": "%Y-%m-%d",
    "DD MMM YYYY": "%d %b %Y",
}


class PackError(ValueError):
    """A request the pack cannot be built from (bad period, unknown section)."""


def _v(value: Any) -> Any:
    return getattr(value, "value", value)


def words(value: Any) -> str:
    """``in_progress`` → ``In progress``; blank → ``—``."""
    text = str(_v(value) or "").replace("_", " ").strip()
    return text[:1].upper() + text[1:] if text else "—"


# ---------------------------------------------------------------------------
# Period and sections (pure)
# ---------------------------------------------------------------------------
def fiscal_quarter_start(day: date, fiscal_start_month: int = 1) -> date:
    """First day of the fiscal quarter ``day`` falls in. Quarters run in threes from the
    fiscal year's first month (January for Pakistani banks, the default)."""
    start0 = (int(fiscal_start_month or 1) - 1) % 12
    back = ((day.month - 1 - start0) % 12) % 3
    month0 = day.month - 1 - back
    year = day.year
    if month0 < 0:
        month0 += 12
        year -= 1
    return date(year, month0 + 1, 1)


def resolve_period(
    start: date | None, end: date | None, *, today: date, fiscal_start_month: int = 1,
) -> tuple[date, date]:
    """The pack's period. Default: the fiscal quarter to date. An end alone runs from the
    start of its quarter. Raises :class:`PackError` for a period that ends in the future
    or starts after it ends."""
    end = end or today
    if end > today:
        raise PackError(f"period_end: {end.isoformat()} is in the future; end the period today or earlier.")
    start = start or fiscal_quarter_start(end, fiscal_start_month)
    if start > end:
        raise PackError("period_start must be on or before period_end.")
    return start, end


def normalise_sections(requested: Iterable[str] | None, keep_order: bool = False) -> list[str]:
    """The sections to include, in pack order (or, with ``keep_order``, in the order
    asked for — a committee's saved order). None means all; an unknown key or an empty
    choice is refused. Duplicates are dropped."""
    if requested is None:
        return list(SECTION_KEYS)
    listed = [str(s).strip() for s in requested if str(s).strip()]
    chosen = set(listed)
    unknown = sorted(chosen - set(SECTION_KEYS))
    if unknown:
        raise PackError(
            f"Unknown section(s): {', '.join(unknown)}. Choose from: {', '.join(SECTION_KEYS)}."
        )
    if not chosen:
        raise PackError("Choose at least one section for the pack.")
    if keep_order:
        return list(dict.fromkeys(listed))
    return [k for k in SECTION_KEYS if k in chosen]


def format_date(value: date | datetime | None, date_format: str = "DD/MM/YYYY") -> str:
    if value is None:
        return "—"
    if isinstance(value, datetime):
        value = value.date()
    return value.strftime(DATE_FORMATS.get(date_format, "%d/%m/%Y"))


def format_datetime(value: datetime | None, date_format: str, tz) -> str:
    """Date in the organisation's format and time in its timezone, e.g. ``12/09/2026 14:05``."""
    if value is None:
        return "—"
    aware = value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)
    local = aware.astimezone(tz) if tz is not None else aware
    return f"{format_date(local.date(), date_format)} {local.strftime('%H:%M')}"


def period_text(start: date, end: date, date_format: str) -> str:
    return f"{format_date(start, date_format)} to {format_date(end, date_format)}"


def pack_title(
    *, committee: str | None, meeting: str | None, start: date, end: date, date_format: str,
) -> str:
    if committee and meeting:
        title = f"Board pack: {committee} — {meeting}"
    elif committee:
        title = f"Board pack: {committee}, {period_text(start, end, date_format)}"
    else:
        title = f"Board pack, {period_text(start, end, date_format)}"
    return title[:255]


def meeting_due(meeting_date: date | None, days_before: int | None, today: date) -> bool:
    """A scheduled meeting is due its pack from ``days_before`` days out until the day."""
    if meeting_date is None or not days_before or days_before < 1:
        return False
    return 0 <= (meeting_date - today).days <= days_before


def pack_is_current(generated_on: date | None, meeting_date: date, days_before: int) -> bool:
    """A pack generated inside the meeting's window counts; one generated before the
    window opened (a draft weeks earlier) is stale, so the scheduler makes a fresh one."""
    return generated_on is not None and generated_on >= meeting_date - timedelta(days=days_before)


def period_instants(start: date, end: date, tz) -> tuple[datetime, datetime]:
    """[start 00:00, day after end 00:00) in the organisation's timezone."""
    return (
        datetime.combine(start, time(0, 0), tzinfo=tz),
        datetime.combine(end + timedelta(days=1), time(0, 0), tzinfo=tz),
    )


def _aware(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)


def in_range(value: datetime | None, start_at: datetime, end_at: datetime) -> bool:
    v = _aware(value)
    return v is not None and start_at <= v < end_at


# ---------------------------------------------------------------------------
# Risk trend and movement (pure)
# ---------------------------------------------------------------------------
TREND_NEW, TREND_UNCHANGED, TREND_HISTORY, TREND_UNKNOWN = "new", "unchanged", "history", "unknown"

TREND_BASIS_TEXT = {
    TREND_NEW: "added in the period",
    TREND_UNCHANGED: "not re-scored since the period began",
    TREND_HISTORY: "score at the start of the period from the record's history",
    TREND_UNKNOWN: "no score recorded before the period began",
}


@dataclass(frozen=True)
class Trend:
    #: up | down | same | new | unknown
    direction: str
    start_score: int | None
    basis: str


def effective_score(inherent: int | None, residual: int | None) -> int | None:
    """Current exposure: the residual once assessed, otherwise the inherent score."""
    return residual if residual is not None else inherent


def _product(a: Any, b: Any) -> int | None:
    try:
        return int(a) * int(b) if a and b else None
    except (TypeError, ValueError):
        return None


def score_from_snapshot(snapshot: Mapping[str, Any] | None) -> int | None:
    """Effective score in a stored version snapshot (likelihood x impact, so a snapshot
    taken before the computed score column was refreshed still reads right)."""
    if not snapshot:
        return None
    inherent = _product(snapshot.get("inherent_likelihood"), snapshot.get("inherent_impact"))
    residual = _product(snapshot.get("residual_likelihood"), snapshot.get("residual_impact"))
    return effective_score(inherent, residual)


def score_at_start(
    snapshot: Mapping[str, Any] | None,
    snapshot_at: datetime | None,
    assessment: Mapping[str, Any] | None,
    assessment_at: datetime | None,
) -> int | None:
    """The effective score when the period began, from the latest version snapshot taken
    before it — unless a residual assessment (recorded in the audit trail, not versioned)
    came after that snapshot, in which case its residual is the score."""
    score = score_from_snapshot(snapshot)
    if assessment and assessment_at is not None and (snapshot_at is None or _aware(assessment_at) > _aware(snapshot_at)):
        residual = _product(assessment.get("residual_likelihood"), assessment.get("residual_impact"))
        if residual is not None:
            return residual
    return score


def risk_trend(
    *,
    current: int | None,
    created_at: datetime | None,
    last_assessed_at: datetime | None,
    period_start_at: datetime,
    snapshot: Mapping[str, Any] | None = None,
    snapshot_at: datetime | None = None,
    assessment: Mapping[str, Any] | None = None,
    assessment_at: datetime | None = None,
) -> Trend:
    """How a risk's score moved since the period began. See the module notes."""
    start = _aware(period_start_at)
    if _aware(created_at) is not None and _aware(created_at) >= start:
        return Trend("new", None, TREND_NEW)
    assessed = _aware(last_assessed_at)
    if assessed is not None and assessed < start:
        return Trend("same", current, TREND_UNCHANGED)
    begin = score_at_start(snapshot, snapshot_at, assessment, assessment_at) if (snapshot or assessment) else None
    if begin is None:
        return Trend("unknown", None, TREND_UNKNOWN)
    if current is None:
        return Trend("unknown", begin, TREND_HISTORY)
    direction = "up" if current > begin else "down" if current < begin else "same"
    return Trend(direction, begin, TREND_HISTORY)


def trend_text(trend: Trend) -> str:
    if trend.direction == "new":
        return "New in period"
    if trend.direction == "unknown":
        return "No earlier score" if trend.start_score is None else f"Was {trend.start_score}"
    if trend.direction == "same":
        return "Unchanged"
    return f"{'Up' if trend.direction == 'up' else 'Down'} from {trend.start_score}"


def classify_risk_movement(rows: Iterable[Any], start_at: datetime, end_at: datetime) -> dict[str, list]:
    """Split risk rows into those added, closed and re-scored in the period.

    * **new** — created in the period;
    * **closed** — status closed and last changed in the period (the dashboard's rule);
    * **re-scored** — scores changed in the period on a risk that existed before it (a
      new risk's first scoring is not a re-score).
    A row may be both new and closed. Rows need ``created_at``, ``updated_at``,
    ``status`` and ``last_assessed_at``.
    """
    new, closed, rescored = [], [], []
    for r in rows:
        created_in = in_range(r.created_at, start_at, end_at)
        if created_in:
            new.append(r)
        if words(r.status).lower() == "closed" and in_range(r.updated_at, start_at, end_at):
            closed.append(r)
        if (not created_in and _aware(r.created_at) is not None and _aware(r.created_at) < start_at
                and in_range(getattr(r, "last_assessed_at", None), start_at, end_at)):
            rescored.append(r)
    return {"new": new, "closed": closed, "rescored": rescored}


# ---------------------------------------------------------------------------
# Period sections (pure summaries over loaded rows)
# ---------------------------------------------------------------------------
def issue_summary(
    rows: Iterable[Any], moves: Mapping[Any, int], today: date, owner_names: Mapping[Any, str] | None = None,
) -> dict:
    """Open issues by severity: how many, how many overdue, how many had their agreed due
    date moved (and how many moves). ``rows`` are open issues; ``moves`` maps issue id to
    its approved date moves; ``owner_names`` maps owner ids to names (the legacy owner
    text is the fallback)."""
    owner_names = owner_names or {}
    by: dict[str, Counter] = {s: Counter() for s in SEVERITIES}
    overdue_rows = []
    regulator = 0
    rows = list(rows)
    for r in rows:
        sev = str(_v(r.severity) or "medium")
        bucket = by.setdefault(sev, Counter())
        bucket["open"] += 1
        n_moves = int(moves.get(r.id, 0) or 0)
        if n_moves:
            bucket["date_moved"] += 1
            bucket["moves"] += n_moves
        if r.due_date is not None and r.due_date < today:
            bucket["overdue"] += 1
            overdue_rows.append((r, (today - r.due_date).days, n_moves))
        if getattr(r, "regulator_related", False):
            regulator += 1
    rank = {s: i for i, s in enumerate(SEVERITIES)}
    overdue_rows.sort(key=lambda t: (rank.get(str(_v(t[0].severity)), 9), -t[1], t[0].reference or ""))
    return {
        "open": len(rows),
        "overdue": sum(b["overdue"] for b in by.values()),
        "date_moved": sum(b["date_moved"] for b in by.values()),
        "regulator_related": regulator,
        "by_severity": [
            {"severity": s, "open": b["open"], "overdue": b["overdue"],
             "date_moved": b["date_moved"], "moves": b["moves"]}
            for s, b in by.items()
        ],
        "overdue_rows": [
            {"reference": r.reference, "title": r.title, "severity": str(_v(r.severity)),
             "owner": owner_names.get(getattr(r, "owner_id", None)) or getattr(r, "owner", "") or "",
             "due": r.due_date, "days_overdue": days, "moves": n}
            for r, days, n in overdue_rows
        ],
    }


def incident_summary(incidents: Iterable[Any], reports: Mapping[Any, Sequence[Any]], now: datetime) -> dict:
    """Incidents in the period: severity mix, near misses, reportable ones and whether
    the regulator was notified on time, and mean time to detect / resolve (hours).

    The on-time rate is taken over notifications that are decided — made (on time or
    late) or already past their deadline — so one still inside its deadline is not
    counted either way."""
    from app.services import incident_clock

    incidents = list(incidents)
    sev = Counter(str(_v(i.severity) or "medium") for i in incidents)
    on_time = late = pending = 0
    rows = []
    mttd, mttr = [], []
    for inc in incidents:
        times = incident_clock.response_times(inc)
        mttd.append(times.mttd_hours)
        mttr.append(times.mttr_hours)
        clock = None
        if inc.is_reportable:
            clock = incident_clock.notification_clock(reports.get(inc.id, ()), now)
            if clock.notified_on_time is True:
                on_time += 1
            elif clock.notified_on_time is False:
                late += 1
            elif clock.notification_deadline is not None:
                pending += 1
            rows.append({
                "reference": inc.reference, "title": inc.title, "severity": str(_v(inc.severity)),
                "regulator": getattr(inc, "regulator", "") or "",
                "detected": inc.detected_at or getattr(inc, "created_at", None),
                "deadline": clock.notification_deadline, "notified": clock.notified_at,
                "on_time": clock.notified_on_time,
            })
    decided = on_time + late
    avg_d, n_d = incident_clock.average(mttd)
    avg_r, n_r = incident_clock.average(mttr)
    return {
        "total": len(incidents),
        "by_severity": {s: sev.get(s, 0) for s in SEVERITIES},
        "near_misses": sum(1 for i in incidents if getattr(i, "near_miss", False)),
        "reportable": sum(1 for i in incidents if i.is_reportable),
        "on_time": on_time, "late": late, "pending": pending,
        "on_time_rate": round(100.0 * on_time / decided, 1) if decided else None,
        "mttd_hours": avg_d, "mttd_n": n_d, "mttr_hours": avg_r, "mttr_n": n_r,
        "rows": rows,
    }


def kri_summary(kris: Iterable[Any], owner_names: Mapping[Any, str] | None = None) -> dict:
    """KRIs by status, and the red then amber ones listed (worst first)."""
    owner_names = owner_names or {}
    counts: Counter = Counter()
    listed = []
    for k in kris:
        status = str(_v(k.status))
        counts[status] += 1
        if status in ("red", "amber"):
            listed.append(k)
    listed.sort(key=lambda k: (0 if str(_v(k.status)) == "red" else 1, k.reference or "", k.name or ""))
    return {
        "green": counts["green"], "amber": counts["amber"], "red": counts["red"],
        "no_data": counts["no_data"],
        "rows": [
            {"id": str(getattr(k, "id", "") or ""), "reference": k.reference or "", "name": k.name,
             "status": str(_v(k.status)),
             "value": _num(k.current_value), "unit": k.unit or "",
             "warning": _num(k.warning_threshold), "limit": _num(k.limit_threshold),
             "lower": _num(getattr(k, "lower_bound", None)), "upper": _num(getattr(k, "upper_bound", None)),
             "direction": str(_v(getattr(k, "direction", "")) or ""),
             "owner": owner_names.get(getattr(k, "owner_id", None)) or k.owner or "",
             "as_of": k.last_measured_date}
            for k in listed
        ],
    }


def third_party_summary(vendors: Iterable[Any], certs: Iterable[Any], today: date) -> dict:
    """Third parties rated critical or high (criticality is the operative rating; the
    inherent tier from the tiering questionnaire is shown beside it) and certifications
    expired or expiring within the warning window. Offboarded third parties are left out
    by the caller."""
    from app.models.vendor import certification_expiry_state

    certs = list(certs)
    lapsing_by_vendor: Counter = Counter()
    cert_rows = []
    for c in certs:
        state = certification_expiry_state(c.expires_on, today)
        if state not in ("expired", "expiring"):
            continue
        lapsing_by_vendor[c.vendor_id] += 1
        cert_rows.append({
            "vendor": getattr(c, "vendor_name", "") or "", "cert_type": c.cert_type, "expires_on": c.expires_on,
            "state": state, "days": (c.expires_on - today).days,
        })
    cert_rows.sort(key=lambda r: (r["expires_on"], r["vendor"]))
    rank = {"critical": 0, "high": 1}
    chosen = [v for v in vendors if str(_v(v.criticality)) in rank]
    chosen.sort(key=lambda v: (rank[str(_v(v.criticality))], v.name.lower()))
    return {
        "critical": sum(1 for v in chosen if str(_v(v.criticality)) == "critical"),
        "high": sum(1 for v in chosen if str(_v(v.criticality)) == "high"),
        "expired_certs": sum(1 for r in cert_rows if r["state"] == "expired"),
        "expiring_certs": sum(1 for r in cert_rows if r["state"] == "expiring"),
        "rows": [
            {"name": v.name, "criticality": str(_v(v.criticality)), "tier": v.inherent_tier or "",
             "status": str(_v(v.status)), "next_review": v.next_review_date,
             "review_overdue": bool(v.next_review_date and v.next_review_date < today),
             "certs_lapsing": lapsing_by_vendor.get(v.id, 0)}
            for v in chosen
        ],
        "certifications": cert_rows,
    }


def failed_tests_summary(rows: Iterable[Any]) -> dict:
    """Control tests conducted in the period that failed or found exceptions. Only
    reviewed tests (and legacy ones recorded before reviews existed) count; a failed test
    still awaiting review is counted separately — it has not changed anything yet."""
    from app.services.control_assurance import COUNTING_REVIEW_STATES

    failed, pending, exceptions = [], 0, 0
    for r in rows:
        counts = str(_v(r.review_status) or "legacy") in COUNTING_REVIEW_STATES
        result = str(_v(r.result))
        if result == "failed":
            if counts:
                failed.append(r)
            else:
                pending += 1
        elif result == "passed_with_exceptions" and counts:
            exceptions += 1
    failed.sort(key=lambda r: (r.conducted_date or date.min, r.control_reference or ""), reverse=True)
    return {
        "failed": [
            {"control": " ".join(p for p in (r.control_reference or "", r.control_name or "") if p),
             "conducted": r.conducted_date, "key": bool(r.is_key), "issue": getattr(r, "issue_reference", "") or ""}
            for r in failed
        ],
        "failed_pending_review": pending,
        "exceptions": exceptions,
    }


def _num(value: Any) -> float | None:
    return float(value) if value is not None else None


def fmt_num(value: Any, unit: str = "") -> str:
    if value is None:
        return "—"
    v = float(value)
    text = str(int(v)) if v.is_integer() else f"{v:,.2f}".rstrip("0").rstrip(".")
    return f"{text} {unit}".strip() if unit else text


def fmt_pct(value: float | None) -> str:
    return "—" if value is None else f"{value:.1f}%".replace(".0%", "%")


def fmt_hours(value: float | None) -> str:
    if value is None:
        return "—"
    if value >= 48:
        return f"{value / 24:.1f} days"
    return f"{value:.1f} h"


# ---------------------------------------------------------------------------
# Section views — what both renderers print (pure)
# ---------------------------------------------------------------------------
@dataclass
class TableView:
    caption: str
    headers: list[str]
    rows: list[list[str]]
    #: Relative column widths; the renderer scales them to the page.
    widths: list[int]
    empty: str = "None."
    #: Total rows before the listing limit, for a "showing 50 of 212" note.
    total: int | None = None


@dataclass
class ChartView:
    """A chart the PDF draws (and the spreadsheet lists as a table of its data).

    * ``heatmap`` — ``data = {"size": n, "cells": {"L,I": count}, "bands": {"L,I": band}}``
    * ``appetite_trend`` — ``data = {"points": [{"label", "within", "elevated", "breach"}]}``
      (counts None where no snapshot stands for the date)
    * ``kri_trend`` — ``data = {"series": [{"name", "unit", "readings": [{"label", "value"}]}]}``
    """

    kind: str
    title: str
    data: dict
    note: str = ""


@dataclass
class SectionView:
    key: str
    title: str
    kpis: list[tuple[str, str]] = field(default_factory=list)
    facts: list[tuple[str, str]] = field(default_factory=list)
    tables: list[TableView] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    #: The narrative written for the committee (phase 4B), printed first.
    commentary: str = ""
    charts: list[ChartView] = field(default_factory=list)


def _limited(rows: list, limit: int | None = LIST_LIMIT) -> tuple[list, int]:
    """The rows to print (all of them when ``limit`` is None) and how many there are."""
    return (rows if limit is None else rows[:limit]), len(rows)


def _view_summary(data: dict, fmt: str, limit: int | None = LIST_LIMIT) -> SectionView:
    view = SectionView("summary", SECTION_TITLES["summary"])
    score = data.get("score")
    band = data.get("band") or ""
    view.kpis = [
        ("Governance health", "No data" if score is None else f"{score} / 100"),
        ("Band", "No data yet" if band == "no_data" else words(band)),
        ("Measures scored", f"{data.get('scored', 0)} of {data.get('total', 0)} ({fmt_pct(data.get('weight_pct'))} of weight)"),
    ]
    view.facts = [(label, value) for label, value in data.get("headlines", [])]
    view.tables.append(TableView(
        "Needs a decision or is overdue",
        ["Item", "Count"],
        [[a["label"], str(a["count"])] for a in data.get("actions", [])],
        [80, 20],
        empty="Nothing is waiting on a decision or past due.",
    ))
    view.tables.append(TableView(
        "What the score is made of",
        ["Measure", "Value", "Weight", "Behind it"],
        [[c["label"], "No data" if c.get("value") is None else fmt_pct(c["value"]),
          f"{round(100 * float(c.get('weight') or 0))}%", c.get("detail") or ""] for c in data.get("components", [])],
        [30, 12, 10, 48],
    ))
    view.notes.append(
        "A measure with nothing to count is left out of the score and the other weights are "
        "scaled up, so the score is only as broad as the coverage shown."
    )
    return view


def _view_appetite(data: dict, fmt: str, limit: int | None = LIST_LIMIT) -> SectionView:
    view = SectionView("appetite", SECTION_TITLES["appetite"])
    view.kpis = [
        ("Risks", str(data.get("total", 0))),
        ("Within appetite", str(data.get("within", 0))),
        ("Elevated", str(data.get("elevated", 0))),
        ("Above tolerance", str(data.get("breach", 0))),
    ]
    view.tables.append(TableView(
        "Each risk against its own category's appetite and tolerance",
        ["Category", "Appetite", "Tolerance", "Risks", "Within", "Elevated", "Breach"],
        [[r["label"], str(r["appetite"]), str(r["tolerance"]), str(r["risks"]), str(r["within"]),
          str(r["elevated"]), str(r["breach"])] for r in data.get("rows", [])],
        [30, 11, 12, 9, 10, 11, 10],
    ))
    heat = data.get("heatmap")
    if heat and heat.get("cells") is not None:
        view.charts.append(ChartView("heatmap", "Where current exposure sits (likelihood x impact)", heat,
                                     note="Risks on the board register, at their residual cell once assessed."))
    trend = data.get("trend") or []
    if trend:
        view.charts.append(ChartView(
            "appetite_trend", "Appetite position over the last four quarter ends",
            {"points": [{"label": format_date(t.get("date"), fmt), "within": t.get("within"),
                         "elevated": t.get("elevated"), "breach": t.get("breach")} for t in trend]},
            note="A quarter end with no period snapshot near it is left blank rather than estimated.",
        ))
    return view


def _view_top_risks(data: dict, fmt: str, limit: int | None = LIST_LIMIT) -> SectionView:
    view = SectionView("top_risks", SECTION_TITLES["top_risks"])
    rows = data.get("rows", [])
    view.tables.append(TableView(
        "Highest current exposures",
        ["Ref", "Risk", "Score", "Severity", "Appetite", "Owner", "Trend"],
        [[r["reference"], r["title"], fmt_num(r.get("score")), words(r.get("severity")),
          words(r.get("appetite_status")), r.get("owner") or "Unassigned", r.get("trend_text") or "—"] for r in rows],
        [10, 34, 8, 11, 12, 14, 15],
        empty="No risks in the register.",
    ))
    view.notes.append(
        "Score is the current exposure: residual where assessed, otherwise inherent. Trend compares it with "
        "the score when the period began — unchanged when the risk was not re-scored since, otherwise from "
        "the risk's version history; a risk with no score recorded before the period says so."
    )
    return view


def _view_movement(data: dict, fmt: str, limit: int | None = LIST_LIMIT) -> SectionView:
    view = SectionView("movement", SECTION_TITLES["movement"])
    counts = data.get("counts", {})
    view.kpis = [
        ("Added", str(counts.get("new", 0))),
        ("Closed", str(counts.get("closed", 0))),
        ("Re-scored", str(counts.get("rescored", 0))),
        ("Re-scored up / down", f"{counts.get('up', 0)} / {counts.get('down', 0)}"),
    ]
    for key, caption, headers, widths in (
        ("new", "Added in the period", ["Ref", "Risk", "Score", "Added"], [12, 58, 10, 20]),
        ("closed", "Closed in the period", ["Ref", "Risk", "Score", "Closed"], [12, 58, 10, 20]),
        ("rescored", "Re-scored in the period", ["Ref", "Risk", "Score", "Change", "Re-scored"], [12, 46, 10, 16, 16]),
    ):
        items, total = _limited(data.get(key, []), limit)
        rows = []
        for r in items:
            base = [r["reference"], r["title"], fmt_num(r.get("score"))]
            if key == "rescored":
                base.append(r.get("trend_text") or "—")
            base.append(format_date(r.get("when"), fmt))
            rows.append(base)
        view.tables.append(TableView(caption, headers, rows, widths, total=total))
    return view


def _view_assurance(data: dict, fmt: str, limit: int | None = LIST_LIMIT) -> SectionView:
    view = SectionView("assurance", SECTION_TITLES["assurance"])
    operating = sum(int(data.get(k, 0)) for k in ("effective", "partially_effective", "ineffective", "not_assessed"))
    assured = int(data.get("effective", 0)) + int(data.get("partially_effective", 0))
    view.kpis = [
        ("Operating controls", str(operating)),
        ("Effective or partial", f"{assured} ({fmt_pct(round(100 * assured / operating, 1) if operating else None)})"),
        ("Failed tests in period", str(len(data.get("failed", [])))),
        ("Tests overdue", str(data.get("tests_overdue", 0))),
    ]
    view.facts = [
        ("Effective", str(data.get("effective", 0))),
        ("Partially effective", str(data.get("partially_effective", 0))),
        ("Ineffective", str(data.get("ineffective", 0))),
        ("Not tested", str(data.get("not_assessed", 0))),
        ("Not in operation (planned or retired)", str(data.get("not_operating", 0))),
        ("Failed their last reviewed test", str(data.get("last_test_failed", 0))),
        ("Tests in the period that found exceptions", str(data.get("exceptions", 0))),
        ("Failed tests in the period awaiting review", str(data.get("failed_pending_review", 0))),
    ]
    items, total = _limited(data.get("failed", []), limit)
    view.tables.append(TableView(
        "Reviewed tests that failed in the period",
        ["Control", "Tested", "Key control", "Issue raised"],
        [[r["control"], format_date(r.get("conducted"), fmt), "Yes" if r.get("key") else "No", r.get("issue") or "—"]
         for r in items],
        [52, 16, 14, 18], total=total,
    ))
    view.notes.append(
        "Effectiveness moves only on a test another person has reviewed; mapped but untested controls count as not tested."
    )
    return view


def _view_compliance(data: dict, fmt: str, limit: int | None = LIST_LIMIT) -> SectionView:
    view = SectionView("compliance", SECTION_TITLES["compliance"])
    view.kpis = [("Clauses assured by a working control", fmt_pct(data.get("overall_assured_pct")))]
    view.tables.append(TableView(
        "Compliance frameworks (maturity and guidance frameworks are not scored)",
        ["Framework", "In scope", "Assured", "Assured %", "Untested", "Failing", "No control", "Gaps", "Compliant"],
        [[r["name"], str(r["applicable"]), str(r["assured"]), fmt_pct(r.get("assured_pct")), str(r["unassessed"]),
          str(r["failing"]), str(r["unmapped"]), str(r["gaps"]), fmt_pct(r.get("compliant_pct"))]
         for r in data.get("rows", [])],
        [22, 9, 10, 10, 11, 9, 10, 7, 12],
        empty="No compliance frameworks installed.",
    ))
    return view


def _view_issues(data: dict, fmt: str, limit: int | None = LIST_LIMIT) -> SectionView:
    view = SectionView("issues", SECTION_TITLES["issues"])
    view.kpis = [
        ("Open", str(data.get("open", 0))),
        ("Overdue", str(data.get("overdue", 0))),
        ("Due date moved", str(data.get("date_moved", 0))),
        ("Regulator-related", str(data.get("regulator_related", 0))),
    ]
    view.tables.append(TableView(
        "Open issues by severity",
        ["Severity", "Open", "Overdue", "Date moved", "Moves"],
        [[words(r["severity"]), str(r["open"]), str(r["overdue"]), str(r["date_moved"]), str(r["moves"])]
         for r in data.get("by_severity", [])],
        [28, 18, 18, 18, 18],
    ))
    items, total = _limited(data.get("overdue_rows", []), limit)
    view.tables.append(TableView(
        "Overdue issues",
        ["Ref", "Issue", "Severity", "Owner", "Due", "Days over", "Moves"],
        [[r["reference"], r["title"], words(r["severity"]), r.get("owner") or "—", format_date(r.get("due"), fmt),
          str(r["days_overdue"]), str(r["moves"])] for r in items],
        [10, 34, 10, 16, 12, 9, 9], total=total,
    ))
    view.notes.append("Date moved counts issues whose agreed due date was changed at least once (the first date set is not a move).")
    return view


def _view_incidents(data: dict, fmt: str, limit: int | None = LIST_LIMIT, tz=None) -> SectionView:
    view = SectionView("incidents", SECTION_TITLES["incidents"])
    decided = int(data.get("on_time", 0)) + int(data.get("late", 0))
    view.kpis = [
        ("Incidents in period", str(data.get("total", 0))),
        ("Regulator-reportable", str(data.get("reportable", 0))),
        ("Notified on time", "—" if data.get("on_time_rate") is None else f"{fmt_pct(data['on_time_rate'])} ({data.get('on_time', 0)} of {decided})"),
        ("Mean time to detect / resolve", f"{fmt_hours(data.get('mttd_hours'))} / {fmt_hours(data.get('mttr_hours'))}"),
    ]
    sev = data.get("by_severity", {})
    view.facts = [
        ("By severity", ", ".join(f"{words(s)} {sev.get(s, 0)}" for s in SEVERITIES)),
        ("Near misses", str(data.get("near_misses", 0))),
        ("Notifications on time / late / still within deadline",
         f"{data.get('on_time', 0)} / {data.get('late', 0)} / {data.get('pending', 0)}"),
        ("Mean time to detect", f"{fmt_hours(data.get('mttd_hours'))} over {data.get('mttd_n', 0)} incident(s) with both times"),
        ("Mean time to resolve", f"{fmt_hours(data.get('mttr_hours'))} over {data.get('mttr_n', 0)} incident(s) with both times"),
    ]
    items, total = _limited(data.get("rows", []), limit)

    def when(value):
        return format_datetime(value, fmt, tz) if isinstance(value, datetime) else format_date(value, fmt)

    view.tables.append(TableView(
        "Regulator-reportable incidents",
        ["Ref", "Incident", "Severity", "Regulator", "Deadline", "Notified", "On time"],
        [[r["reference"], r["title"], words(r["severity"]), r.get("regulator") or "—", when(r.get("deadline")),
          when(r.get("notified")),
          "Yes" if r.get("on_time") is True else "No" if r.get("on_time") is False else "Pending"]
         for r in items],
        [9, 29, 10, 10, 15, 15, 9], total=total,
    ))
    view.notes.append(
        "Incidents detected in the period (logged, when no detection time was recorded). The on-time rate counts "
        "notifications made or already past their deadline."
    )
    return view


def _view_kris(data: dict, fmt: str, limit: int | None = LIST_LIMIT) -> SectionView:
    view = SectionView("kris", SECTION_TITLES["kris"])
    view.kpis = [
        ("Red", str(data.get("red", 0))), ("Amber", str(data.get("amber", 0))),
        ("Green", str(data.get("green", 0))), ("No data", str(data.get("no_data", 0))),
    ]
    rows = []
    listed, total = _limited(data.get("rows", []), limit)
    for r in listed:
        if r.get("direction") == "within_range":
            limits = f"range {fmt_num(r.get('lower'))}–{fmt_num(r.get('upper'))}"
        else:
            limits = f"warning {fmt_num(r.get('warning'))}, limit {fmt_num(r.get('limit'))}"
        rows.append([r["reference"] or "—", r["name"], words(r["status"]), fmt_num(r.get("value"), r.get("unit", "")),
                     limits, r.get("owner") or "—", format_date(r.get("as_of"), fmt)])
    view.tables.append(TableView(
        "KRIs at red or amber",
        ["Ref", "KRI", "Status", "Value", "Thresholds", "Owner", "As of"],
        rows, [9, 25, 8, 11, 20, 13, 14], total=total,
        empty="No KRI is at red or amber.",
    ))
    series = []
    for r in data.get("rows", []):
        readings = (data.get("trend") or {}).get(str(r.get("id") or ""), [])
        if len(readings) >= 2:
            series.append({"name": " ".join(p for p in (r.get("reference") or "", r["name"]) if p),
                           "unit": r.get("unit", ""),
                           "readings": [{"label": format_date(x.get("as_of"), fmt), "value": x.get("value")}
                                        for x in readings]})
    if series:
        view.charts.append(ChartView("kri_trend", "Recent readings of the KRIs at red or amber",
                                     {"series": series[:KRI_TREND_CHARTS]}))
    return view


def _view_third_parties(data: dict, fmt: str, limit: int | None = LIST_LIMIT) -> SectionView:
    view = SectionView("third_parties", SECTION_TITLES["third_parties"])
    view.kpis = [
        ("Critical", str(data.get("critical", 0))), ("High", str(data.get("high", 0))),
        ("Certifications expired", str(data.get("expired_certs", 0))),
        ("Certifications expiring (60 days)", str(data.get("expiring_certs", 0))),
    ]
    items, total = _limited(data.get("rows", []), limit)
    view.tables.append(TableView(
        "Third parties rated critical or high",
        ["Third party", "Criticality", "Inherent tier", "Status", "Next review", "Certs lapsing"],
        [[r["name"], words(r["criticality"]), words(r.get("tier")), words(r["status"]),
          format_date(r.get("next_review"), fmt) + (" (overdue)" if r.get("review_overdue") else ""),
          str(r.get("certs_lapsing", 0))] for r in items],
        [30, 13, 13, 12, 18, 14], total=total,
    ))
    items, total = _limited(data.get("certifications", []), limit)
    view.tables.append(TableView(
        "Certifications expired or expiring",
        ["Third party", "Certification", "Expires", "State"],
        [[r["vendor"], r["cert_type"], format_date(r.get("expires_on"), fmt),
          "Expired" if r["state"] == "expired" else f"Expires in {r['days']} day(s)"] for r in items],
        [36, 24, 18, 22], total=total,
    ))
    return view


_VIEW_BUILDERS = {
    "summary": _view_summary,
    "appetite": _view_appetite,
    "top_risks": _view_top_risks,
    "movement": _view_movement,
    "assurance": _view_assurance,
    "compliance": _view_compliance,
    "issues": _view_issues,
    "kris": _view_kris,
    "third_parties": _view_third_parties,
}


def section_views(pack: Mapping[str, Any], limit: int | None = LIST_LIMIT) -> list[SectionView]:
    """The printable content of each included section, in pack order. ``limit`` caps the
    rows listed per table (the PDF); None lists everything (the spreadsheet)."""
    cover = pack.get("cover", {})
    fmt = cover.get("date_format", "DD/MM/YYYY")
    tz = cover.get("tz")
    commentary = pack.get("commentary") or {}
    position_notes = pack.get("position_notes") or {}
    out = []
    for key in pack.get("sections", []):
        data = pack.get(key)
        if data is None:
            continue
        if key == "incidents":
            view = _view_incidents(data, fmt, limit, tz)
        else:
            view = _VIEW_BUILDERS[key](data, fmt, limit)
        view.commentary = str(commentary.get(key) or "").strip()
        view.notes = list(position_notes.get(key, [])) + view.notes
        out.append(view)
    return out


def cover_facts(pack: Mapping[str, Any]) -> list[tuple[str, str]]:
    c = pack.get("cover", {})
    fmt = c.get("date_format", "DD/MM/YYYY")
    facts = [("Organisation", c.get("organisation") or "—")]
    if c.get("committee"):
        facts.append(("Committee", " ".join(p for p in (c.get("committee_reference", ""), c["committee"]) if p)))
    if c.get("meeting"):
        meeting = " ".join(p for p in (c.get("meeting_reference", ""), c["meeting"]) if p)
        when = format_date(c.get("meeting_date"), fmt) if c.get("meeting_date") else ""
        facts.append(("Meeting", f"{meeting}, {when}" if when else meeting))
    basis = pack.get("basis") or {}
    if basis.get("position") == "snapshot" and basis.get("snapshot_as_of"):
        position = (f"{format_date(basis['snapshot_as_of'], fmt)} (the period snapshot nearest the period end)")
    elif basis.get("position") == "live" and basis.get("reason") == "no_snapshot":
        position = (f"{format_date(c.get('as_of'), fmt)} (no period snapshot is held near the period end, "
                    "so position figures are as at generation)")
    else:
        position = format_date(c.get("as_of"), fmt)
    facts += [
        ("Period", period_text(c["period_start"], c["period_end"], fmt)),
        ("Position as at", position),
        ("Generated", format_datetime(c.get("generated_at"), fmt, c.get("tz"))
         + (f" ({c.get('timezone')})" if c.get("timezone") else "")),
        ("Generated by", c.get("generated_by") or "—"),
        ("Sections", ", ".join(SECTION_TITLES[k] for k in pack.get("sections", []))),
    ]
    return facts


def pack_subtitle(pack: Mapping[str, Any]) -> str:
    c = pack.get("cover", {})
    fmt = c.get("date_format", "DD/MM/YYYY")
    parts = [p for p in (c.get("committee"), c.get("meeting")) if p]
    parts.append(period_text(c["period_start"], c["period_end"], fmt))
    return " · ".join(parts)


# ---------------------------------------------------------------------------
# XLSX (pure)
# ---------------------------------------------------------------------------
_SHEET_BAD = set('[]:*?/\\')


def sheet_name(title: str, used: set[str]) -> str:
    """An Excel-safe, unique sheet name (31 characters, none of []:*?/\\)."""
    base = "".join(ch for ch in title if ch not in _SHEET_BAD).strip()[:31] or "Sheet"
    name, n = base, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def to_xlsx(pack: Mapping[str, Any]) -> bytes:
    """The pack as a workbook: a cover sheet, then one sheet per section."""
    from fastapi import HTTPException, status

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:  # pragma: no cover
        raise HTTPException(
            status_code=status.HTTP_501_NOT_IMPLEMENTED,
            detail="Excel export requires the 'openpyxl' package on the server.",
        ) from exc

    bold = Font(bold=True)
    head_font = Font(bold=True, color="FFFFFF")
    colour = brand_colour((pack.get("cover") or {}).get("branding"))
    head_fill = PatternFill("solid", fgColor=colour.lstrip("#").upper())
    wrap = Alignment(wrap_text=True, vertical="top")
    used: set[str] = set()

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name("Cover", used)
    ws.append([pack.get("title", "Board pack")])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    for label, value in cover_facts(pack):
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])
    ws.append(["Note", "Position figures are as at the date shown; movement and incidents cover the period."])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90

    for view in section_views(pack, limit=None):
        ws = wb.create_sheet(sheet_name(view.title, used))
        ws.append([view.title])
        ws.cell(row=1, column=1).font = Font(bold=True, size=14)
        if view.commentary:
            ws.append(["Commentary", view.commentary])
            ws.cell(row=ws.max_row, column=1).font = bold
            ws.cell(row=ws.max_row, column=2).alignment = wrap
            ws.append([])
        for label, value in list(view.kpis) + list(view.facts):
            ws.append([label, value])
            ws.cell(row=ws.max_row, column=1).font = bold
        widest = 2
        for table in view.tables:
            ws.append([])
            ws.append([table.caption])
            ws.cell(row=ws.max_row, column=1).font = bold
            ws.append(list(table.headers))
            for cell in ws[ws.max_row]:
                cell.font = head_font
                cell.fill = head_fill
            if not table.rows:
                ws.append([table.empty])
            for row in table.rows:
                ws.append(list(row))
                for cell in ws[ws.max_row]:
                    cell.alignment = wrap
            if table.total is not None and table.total > len(table.rows):
                ws.append([f"Showing {len(table.rows)} of {table.total}."])
            widest = max(widest, len(table.headers))
        for chart in view.charts:
            caption, headers, rows = chart_table(chart)
            ws.append([])
            ws.append([caption])
            ws.cell(row=ws.max_row, column=1).font = bold
            ws.append(list(headers))
            for cell in ws[ws.max_row]:
                cell.font = head_font
                cell.fill = head_fill
            for row in rows:
                ws.append(list(row))
            widest = max(widest, len(headers))
        for note in view.notes:
            ws.append([])
            ws.append([note])
        ws.column_dimensions["A"].width = 30
        for index in range(2, widest + 1):
            ws.column_dimensions[get_column_letter(index)].width = 18 if index != 2 else 44
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def chart_table(chart: ChartView) -> tuple[str, list[str], list[list]]:
    """A chart's data as a caption, headers and rows (for the spreadsheet). Pure."""
    if chart.kind == "heatmap":
        size = int(chart.data.get("size") or 5)
        cells = chart.data.get("cells") or {}
        headers = ["Likelihood \\ impact"] + [str(i) for i in range(1, size + 1)]
        rows = [[str(like)] + [int(cells.get(f"{like},{imp}", 0)) for imp in range(1, size + 1)]
                for like in range(size, 0, -1)]
        return f"{chart.title} (risks per cell)", headers, rows
    if chart.kind == "appetite_trend":
        rows = [[p.get("label"), *(("No snapshot" if p.get(k) is None else p.get(k)) for k in ("within", "elevated", "breach"))]
                for p in chart.data.get("points", [])]
        return chart.title, ["As at", "Within appetite", "Elevated", "Above tolerance"], rows
    rows = []
    for s in chart.data.get("series", []):
        for r in s.get("readings", []):
            rows.append([s.get("name"), r.get("label"), r.get("value"), s.get("unit", "")])
    return chart.title, ["KRI", "As of", "Value", "Unit"], rows


DEFAULT_CLASSIFICATION = "Confidential"


def brand_colour(branding: Mapping[str, Any] | None) -> str:
    """The pack's primary colour: the organisation's when it is a valid ``#RRGGBB``, else
    the product colour. Pure."""
    import re

    from app.services.pdf_report import PRIMARY

    value = str((branding or {}).get("primary_colour") or "").strip()
    return value if re.fullmatch(r"#[0-9a-fA-F]{6}", value) else PRIMARY


def to_pdf(pack: Mapping[str, Any]) -> bytes:
    from app.services import pdf_report

    cover = pack.get("cover", {})
    branding = dict(cover.get("branding") or {})
    return pdf_report.board_pack_pdf(
        title=pack.get("title", "Board pack"), subtitle=pack_subtitle(pack),
        cover=cover_facts(pack), views=section_views(pack),
        org_name=cover.get("organisation") or "Organisation",
        colour=brand_colour(branding),
        cover_title=str(branding.get("cover_title") or ""),
        classification=str(branding.get("classification") or DEFAULT_CLASSIFICATION),
        logo_path=branding.get("logo_path") or None,
        draft=str(pack.get("review_state") or "") in (DRAFT, REVIEWED),
    )


def file_stem(title: str, period_end: date) -> str:
    slug = "".join(ch if ch.isalnum() else "-" for ch in (title or "board-pack").lower())
    slug = "-".join(p for p in slug.split("-") if p)[:70] or "board-pack"
    return f"{slug}-{period_end.isoformat()}"


def pdf_text(value: Any) -> str:
    """Escape user text for ReportLab's paragraph markup."""
    return escape(str(value if value not in (None, "") else "—"))


# ===========================================================================
# Database: load one organisation's figures, store the files, write the row
# ===========================================================================
@dataclass
class OrgContext:
    tenant_id: uuid.UUID
    name: str
    tz: Any
    timezone: str
    date_format: str
    fiscal_start_month: int
    today: date
    now: datetime


@dataclass(frozen=True)
class _TenantActor:
    """The caller the dashboard overview sees when the scheduler builds a pack: the
    organisation, with no person behind it (the platform reads everything the
    organisation's row-level security scope holds)."""

    tenant_id: uuid.UUID
    id: None = None
    email: str = "system@nexusline"
    full_name: str = SCHEDULER_NAME
    is_active: bool = True
    permission_codes: tuple = ()

    @classmethod
    def for_tenant(cls, tenant_id) -> "_TenantActor":
        from app.core.permissions import ALL_PERMISSIONS

        return cls(tenant_id=uuid.UUID(str(tenant_id)), permission_codes=tuple(ALL_PERMISSIONS))


async def org_context(db, tenant_id) -> OrgContext:
    """Name, timezone, date format, fiscal year and today for the organisation."""
    from sqlalchemy import select

    from app.api.v1.tenant_settings import get_or_create_settings
    from app.models.tenant import Tenant
    from app.services import incident_clock

    row = await get_or_create_settings(db, tenant_id)
    name = await db.scalar(select(Tenant.name).where(Tenant.id == tenant_id)) or "Organisation"
    tz = incident_clock.zone(getattr(row, "timezone", None))
    now = datetime.now(timezone.utc)
    return OrgContext(
        tenant_id=uuid.UUID(str(tenant_id)), name=name, tz=tz,
        timezone=getattr(row, "timezone", "") or incident_clock.DEFAULT_TIMEZONE,
        date_format=getattr(row, "date_format", "") or "DD/MM/YYYY",
        fiscal_start_month=int(getattr(row, "fiscal_year_start_month", 1) or 1),
        today=now.astimezone(tz).date(), now=now,
    )


async def module_enabled(db, key: str) -> bool:
    """The module is licensed on this installation and switched on for the organisation
    (the same two checks ``require_module`` makes for a signed-in request)."""
    from sqlalchemy import select

    from app.models.settings import TenantSettings
    from app.services import modules

    if not modules.is_enabled(key):
        return False
    chosen = await db.scalar(select(TenantSettings.enabled_modules))
    return not isinstance(chosen, list) or key in chosen


async def dashboard_overview(db, tenant_id, days: int, viewer: Any = None):
    """The dashboard's own overview for this organisation — called, not re-derived, so
    the pack and the dashboard can never disagree. ``viewer`` is the person generating
    the pack (the scheduler's system actor when None). Any further query parameter the
    endpoint grows is passed its declared default."""
    import inspect

    from pydantic.fields import FieldInfo

    from app.api.v1 import dashboard

    fn = dashboard.get_overview
    kwargs: dict[str, Any] = {}
    for name, param in inspect.signature(fn).parameters.items():
        if name == "db":
            kwargs[name] = db
        elif name == "user":
            kwargs[name] = viewer if viewer is not None else _TenantActor.for_tenant(tenant_id)
        elif name == "days":
            kwargs[name] = days
        elif param.default is not inspect.Parameter.empty:
            default = param.default
            kwargs[name] = default.default if isinstance(default, FieldInfo) else default
    return await fn(**kwargs)


def summary_from_overview(overview) -> dict:
    health = overview.health
    coverage = health.coverage
    scored = health.band != "no_data"
    a = overview.assurance
    operating = a.effective + a.partially_effective + a.ineffective + a.not_assessed
    p = overview.posture
    return {
        "score": health.score if scored else None,
        "band": health.band,
        "scored": coverage.scored if coverage else 0,
        "total": coverage.total if coverage else len(health.components),
        "weight_pct": coverage.weight_pct if coverage else None,
        "components": [
            {"label": c.label, "value": c.value, "weight": c.weight, "detail": c.detail}
            for c in health.components
        ],
        "headlines": [
            ("Risks above tolerance", f"{p.breach} of {p.total_risks}"),
            ("Operating controls effective or partially effective",
             f"{a.effective + a.partially_effective} of {operating}"),
            ("Clauses assured by a working control", fmt_pct(overview.compliance.overall_assured_pct)),
            ("Open incidents", f"{overview.incidents.open} ({overview.incidents.reportable_open} regulator-reportable)"),
            ("KRIs at red", str(overview.kris.red)),
        ],
        "actions": [{"label": x.label, "count": x.count} for x in overview.actions],
    }


def appetite_from_overview(overview) -> dict:
    p = overview.posture
    rows = [
        {"label": c.label, "appetite": c.appetite_score, "tolerance": c.tolerance_score, "risks": c.risks,
         "within": c.within_appetite, "elevated": c.elevated, "breach": c.breach}
        for c in (p.by_category or [])
    ]
    if not rows:
        rows = [{"label": "All categories (organisation appetite)", "appetite": p.appetite_score,
                 "tolerance": p.tolerance_score, "risks": p.total_risks, "within": p.within_appetite,
                 "elevated": p.elevated, "breach": p.breach}]
    return {"appetite": p.appetite_score, "tolerance": p.tolerance_score, "total": p.total_risks,
            "within": p.within_appetite, "elevated": p.elevated, "breach": p.breach, "rows": rows}


def compliance_from_overview(overview) -> dict:
    return {
        "overall_assured_pct": overview.compliance.overall_assured_pct,
        "rows": [
            {"name": f.name, "applicable": f.applicable, "assured": f.assured,
             "assured_pct": round(100.0 * f.assured / f.applicable, 1) if f.applicable else None,
             "unassessed": f.unassessed, "failing": f.failing, "unmapped": f.unmapped, "gaps": f.gaps,
             "compliant_pct": f.compliant_pct}
            for f in overview.compliance.frameworks
        ],
    }


async def _history_before(db, risk_ids: Sequence, start_at: datetime) -> tuple[dict, dict]:
    """For each risk, the latest version snapshot and the latest residual assessment
    recorded at or before ``start_at``: ({id: (snapshot, at)}, {id: (changes, at)})."""
    from sqlalchemy import select

    from app.models.audit import AuditLog
    from app.models.version import RecordVersion

    ids = list(dict.fromkeys(risk_ids))
    if not ids:
        return {}, {}
    snaps = {
        rid: (snap or {}, at)
        for rid, snap, at in (await db.execute(
            select(RecordVersion.entity_id, RecordVersion.snapshot, RecordVersion.created_at)
            .where(RecordVersion.entity_type == "risk", RecordVersion.entity_id.in_(ids),
                   RecordVersion.created_at <= start_at)
            .distinct(RecordVersion.entity_id)
            .order_by(RecordVersion.entity_id, RecordVersion.created_at.desc(), RecordVersion.version_no.desc())
        )).all()
    }
    assessed = {
        rid: (changes or {}, at)
        for rid, changes, at in (await db.execute(
            select(AuditLog.entity_id, AuditLog.changes, AuditLog.created_at)
            .where(AuditLog.entity_type == "risk", AuditLog.action == "assess", AuditLog.entity_id.in_(ids),
                   AuditLog.created_at <= start_at)
            .distinct(AuditLog.entity_id)
            .order_by(AuditLog.entity_id, AuditLog.created_at.desc())
        )).all()
    }
    return snaps, assessed


async def _trends(db, rows: Sequence[Any], start_at: datetime, current: Mapping[Any, int | None]) -> dict:
    """Trend per risk row (rows need id, created_at, last_assessed_at)."""
    need = [
        r.id for r in rows
        if not (_aware(r.created_at) is not None and _aware(r.created_at) >= start_at)
        and not (_aware(r.last_assessed_at) is not None and _aware(r.last_assessed_at) < start_at)
    ]
    snaps, assessed = await _history_before(db, need, start_at)
    out = {}
    for r in rows:
        snap = snaps.get(r.id)
        assessment = assessed.get(r.id)
        out[r.id] = risk_trend(
            current=current.get(r.id), created_at=r.created_at, last_assessed_at=r.last_assessed_at,
            period_start_at=start_at,
            snapshot=snap[0] if snap else None, snapshot_at=snap[1] if snap else None,
            assessment=assessment[0] if assessment else None, assessment_at=assessment[1] if assessment else None,
        )
    return out


async def _top_risks(db, overview, start_at: datetime) -> dict:
    from sqlalchemy import select

    from app.models.risk import Risk

    top = list(overview.posture.top_risks)
    meta = {}
    if top:
        meta = {r.id: r for r in (await db.execute(
            select(Risk.id, Risk.created_at, Risk.last_assessed_at).where(Risk.id.in_([t.id for t in top]))
        )).all()}
    trends = await _trends(db, list(meta.values()), start_at, {t.id: t.score for t in top})
    rows = []
    for t in top:
        trend = trends.get(t.id) or Trend("unknown", None, TREND_UNKNOWN)
        rows.append({
            "reference": t.reference, "title": t.title, "score": t.score, "severity": t.severity,
            "appetite_status": t.appetite_status, "owner": t.owner, "status": t.status,
            "trend": trend.direction, "start_score": trend.start_score, "trend_basis": trend.basis,
            "trend_text": trend_text(trend),
        })
    return {"rows": rows}


async def _movement(db, start_at: datetime, end_at: datetime) -> dict:
    from sqlalchemy import and_, or_, select

    from app.models.audit import AuditLog
    from app.models.enums import RiskStatus
    from app.models.risk import Risk

    # Risks re-assessed in the period according to the trail (a residual assessment, or
    # an edit that changed or re-confirmed the scores), as well as by their latest stamp.
    reassessed_ids = select(AuditLog.entity_id).where(
        AuditLog.entity_type == "risk", AuditLog.created_at >= start_at, AuditLog.created_at < end_at,
        or_(AuditLog.action == "assess", and_(AuditLog.action == "update", AuditLog.changes.has_key("assessed"))),
    )
    rows = (await db.execute(
        select(Risk.id, Risk.reference, Risk.title, Risk.status, Risk.created_at, Risk.updated_at,
               Risk.last_assessed_at, Risk.inherent_score, Risk.residual_score)
        .where(
            Risk.deleted.is_(False),
            or_(
                and_(Risk.created_at >= start_at, Risk.created_at < end_at),
                and_(Risk.status == RiskStatus.closed, Risk.updated_at >= start_at, Risk.updated_at < end_at),
                and_(Risk.last_assessed_at >= start_at, Risk.last_assessed_at < end_at),
                Risk.id.in_(reassessed_ids),
            ),
        )
    )).all()
    from_trail = set((await db.scalars(reassessed_ids)).all())
    moved = classify_risk_movement(rows, start_at, end_at)
    already = {r.id for r in moved["rescored"]}
    for r in rows:
        if r.id in from_trail and r.id not in already and _aware(r.created_at) < start_at:
            moved["rescored"].append(r)
            already.add(r.id)
    current = {r.id: effective_score(r.inherent_score, r.residual_score) for r in rows}
    trends = await _trends(db, moved["rescored"], start_at, current)

    def row(r, when, trend=None) -> dict:
        out = {"reference": r.reference, "title": r.title, "score": current.get(r.id), "when": when}
        if trend is not None:
            out["trend_text"] = trend_text(trend)
        return out

    def by_ref(items):
        return sorted(items, key=lambda r: r.reference or "")

    rescored_rows = [
        row(r, _aware(r.last_assessed_at) or _aware(r.updated_at), trends.get(r.id)) for r in by_ref(moved["rescored"])
    ]
    ups = sum(1 for r in moved["rescored"] if trends.get(r.id) and trends[r.id].direction == "up")
    downs = sum(1 for r in moved["rescored"] if trends.get(r.id) and trends[r.id].direction == "down")
    return {
        "new": [row(r, r.created_at) for r in by_ref(moved["new"])],
        "closed": [row(r, r.updated_at) for r in by_ref(moved["closed"])],
        "rescored": rescored_rows,
        "counts": {"new": len(moved["new"]), "closed": len(moved["closed"]),
                   "rescored": len(moved["rescored"]), "up": ups, "down": downs},
    }


async def _assurance(db, overview, start: date, end: date) -> dict:
    from types import SimpleNamespace

    from sqlalchemy import select

    from app.models.control import Control, ControlAudit
    from app.models.enums import TestResult
    from app.models.issue import Issue

    a = overview.assurance
    tests = (await db.execute(
        select(ControlAudit.id, ControlAudit.conducted_date, ControlAudit.result, ControlAudit.review_status,
               ControlAudit.raised_issue_id, Control.reference.label("control_reference"),
               Control.name.label("control_name"), Control.is_key)
        .join(Control, Control.id == ControlAudit.control_id)
        .where(Control.deleted.is_(False), ControlAudit.conducted_date >= start, ControlAudit.conducted_date <= end,
               ControlAudit.result.in_((TestResult.failed, TestResult.passed_with_exceptions)))
    )).all()
    issue_ids = {t.raised_issue_id for t in tests if t.raised_issue_id}
    issue_refs = {}
    if issue_ids:
        issue_refs = dict((await db.execute(select(Issue.id, Issue.reference).where(Issue.id.in_(issue_ids)))).all())
    rows = [SimpleNamespace(**t._asdict(), issue_reference=issue_refs.get(t.raised_issue_id, "")) for t in tests]
    failed = failed_tests_summary(rows)
    return {
        "total": a.total, "effective": a.effective, "partially_effective": a.partially_effective,
        "ineffective": a.ineffective, "not_assessed": a.not_assessed, "not_operating": a.not_operating,
        "tests_overdue": a.tests_overdue, "last_test_failed": a.last_test_failed, **failed,
    }


async def _issues(db, today: date) -> dict:
    from sqlalchemy import func, select

    from app.models.issue import Issue, IssueDueDateChange, IssueStatus2
    from app.services import master_data

    closed = [IssueStatus2(s) for s in CLOSED_ISSUE_STATES]
    open_issue = (Issue.deleted.is_(False), Issue.status.not_in(closed))
    rows = (await db.execute(
        select(Issue.id, Issue.reference, Issue.title, Issue.severity, Issue.due_date, Issue.owner,
               Issue.owner_id, Issue.regulator_related).where(*open_issue)
    )).all()
    moves = dict((await db.execute(
        select(IssueDueDateChange.issue_id, func.count())
        .join(Issue, Issue.id == IssueDueDateChange.issue_id)
        .where(*open_issue, IssueDueDateChange.status == "approved", IssueDueDateChange.old_due_date.is_not(None))
        .group_by(IssueDueDateChange.issue_id)
    )).all())
    overdue_owner_ids = [r.owner_id for r in rows if r.due_date is not None and r.due_date < today]
    people = await master_data.users_by_id(db, overdue_owner_ids)
    names = {uid: (ref.full_name or ref.email) for uid, ref in people.items()}
    return issue_summary(rows, moves, today, names)


async def _incidents(db, start_at: datetime, end_at: datetime, now: datetime) -> dict:
    from sqlalchemy import func, select

    from app.models.incident import Incident, RegulatoryReport

    anchor = func.coalesce(Incident.detected_at, Incident.created_at)
    rows = (await db.execute(
        select(Incident.id, Incident.reference, Incident.title, Incident.severity, Incident.is_reportable,
               Incident.regulator, Incident.near_miss, Incident.occurred_at, Incident.detected_at,
               Incident.contained_at, Incident.resolved_at, Incident.created_at)
        .where(Incident.deleted.is_(False), anchor >= start_at, anchor < end_at)
        .order_by(anchor)
    )).all()
    reportable = [r.id for r in rows if r.is_reportable]
    reports: dict = {}
    if reportable:
        for rep in (await db.scalars(select(RegulatoryReport).where(RegulatoryReport.incident_id.in_(reportable)))).all():
            reports.setdefault(rep.incident_id, []).append(rep)
    return incident_summary(rows, reports, now)


async def _kris(db) -> dict:
    from sqlalchemy import select
    from sqlalchemy.orm import noload

    from app.models.operational_risk import KeyRiskIndicator
    from app.services import master_data

    kris = (await db.scalars(
        select(KeyRiskIndicator).where(KeyRiskIndicator.deleted.is_(False)).options(noload("*"))
    )).all()
    flagged = [k.owner_id for k in kris if str(_v(k.status)) in ("red", "amber")]
    people = await master_data.users_by_id(db, flagged)
    return kri_summary(kris, {uid: (ref.full_name or ref.email) for uid, ref in people.items()})


async def _third_parties(db, today: date) -> dict:
    from sqlalchemy import select

    from app.models.enums import VendorStatus
    from app.models.vendor import CERT_EXPIRY_WARNING_DAYS, Vendor, VendorCertification

    live = (Vendor.deleted.is_(False), Vendor.status != VendorStatus.offboarded)
    vendors = (await db.execute(
        select(Vendor.id, Vendor.name, Vendor.criticality, Vendor.inherent_tier, Vendor.status,
               Vendor.next_review_date).where(*live)
    )).all()
    certs = (await db.execute(
        select(VendorCertification.vendor_id, VendorCertification.cert_type, VendorCertification.expires_on,
               Vendor.name.label("vendor_name"))
        .join(Vendor, Vendor.id == VendorCertification.vendor_id)
        .where(*live, VendorCertification.expires_on.is_not(None),
               VendorCertification.expires_on <= today + timedelta(days=CERT_EXPIRY_WARNING_DAYS))
    )).all()
    return third_party_summary(vendors, certs, today)


async def build_pack(
    db, org: OrgContext, *, period_start: date, period_end: date, sections: Sequence[str],
    title: str, committee: Any = None, meeting: Any = None, generated_by: str = "", viewer: Any = None,
) -> dict:
    """Load every included section's figures for the organisation."""
    start_at, end_at = period_instants(period_start, period_end, org.tz)
    pack: dict[str, Any] = {
        "title": title,
        "sections": list(sections),
        "cover": {
            "organisation": org.name,
            "committee": getattr(committee, "name", "") or "",
            "committee_reference": getattr(committee, "reference", "") or "",
            "meeting": getattr(meeting, "title", "") or "",
            "meeting_reference": getattr(meeting, "reference", "") or "",
            "meeting_date": getattr(meeting, "meeting_date", None),
            "period_start": period_start, "period_end": period_end,
            "as_of": org.today, "generated_at": org.now, "generated_by": generated_by,
            "timezone": org.timezone, "tz": org.tz, "date_format": org.date_format,
        },
    }
    overview = None
    if OVERVIEW_SECTIONS & set(sections):
        days = max(7, min(366, (period_end - period_start).days + 1))
        overview = await dashboard_overview(db, org.tenant_id, days, viewer)
    if "summary" in sections:
        pack["summary"] = summary_from_overview(overview)
    if "appetite" in sections:
        pack["appetite"] = appetite_from_overview(overview)
    if "top_risks" in sections:
        pack["top_risks"] = await _top_risks(db, overview, start_at)
    if "movement" in sections:
        pack["movement"] = await _movement(db, start_at, end_at)
    if "assurance" in sections:
        pack["assurance"] = await _assurance(db, overview, period_start, period_end)
    if "compliance" in sections:
        pack["compliance"] = compliance_from_overview(overview)
    if "issues" in sections:
        pack["issues"] = await _issues(db, org.today)
    if "incidents" in sections:
        pack["incidents"] = await _incidents(db, start_at, end_at, org.now)
    if "kris" in sections:
        pack["kris"] = await _kris(db)
    if "third_parties" in sections:
        pack["third_parties"] = await _third_parties(db, org.today)
    # Phase 4B: branding, charts, and position figures for a past period.
    pack["cover"]["branding"] = await load_branding(db)
    pack["basis"] = {"position": "live"}
    snapshot = None
    if period_end < org.today:
        snapshot = await snapshot_for(db, period_end)
        if snapshot is None:
            pack["basis"] = {"position": "live", "reason": "no_snapshot"}
        else:
            as_of, figures = snapshot
            apply_snapshot(pack, figures, as_of, date_format=org.date_format)
            pack["basis"] = {"position": "snapshot", "snapshot_as_of": as_of, "period_end": period_end}
    if "appetite" in sections:
        await add_appetite_charts(db, pack, org, period_end, snapshot)
    if "kris" in sections:
        await add_kri_trend(db, pack, period_end)
    return pack


async def store_file(
    db, tenant_id, *, data: bytes, filename: str, content_type: str, entity_type: str, entity_id,
    title: str, uploaded_by: str,
):
    """Write generated bytes through the same storage seam uploads use and record the
    ``StoredFile`` row (so the file also shows under the record's Files)."""
    from starlette.datastructures import Headers, UploadFile

    from app.models.collab import StoredFile
    from app.services import storage

    upload = UploadFile(io.BytesIO(data), size=len(data), filename=filename,
                        headers=Headers({"content-type": content_type}))
    blob = await storage.save_upload(tenant_id, upload)
    sf = StoredFile(
        tenant_id=tenant_id, entity_type=entity_type, entity_id=entity_id, title=title[:255],
        filename=blob.filename, content_type=blob.content_type, size_bytes=blob.size_bytes,
        sha256=blob.sha256, storage_key=blob.storage_key, uploaded_by_email=(uploaded_by or "")[:255],
    )
    db.add(sf)
    await db.flush()
    return sf


def attach_target(pack_id, committee: Any = None, meeting: Any = None) -> tuple[str, Any]:
    """Where the pack's files are filed: the meeting it was prepared for, else the
    committee, else the pack itself."""
    if meeting is not None:
        return "committee_meeting", meeting.id
    if committee is not None:
        return "committee", committee.id
    return "board_pack", pack_id


def error_text(exc: BaseException) -> str:
    detail = getattr(exc, "detail", None)
    text = detail if isinstance(detail, str) else str(exc) or exc.__class__.__name__
    return f"{text}"[:1000]


def actor_label(actor: Any) -> str:
    if actor is None:
        return SCHEDULER_NAME
    return getattr(actor, "full_name", "") or getattr(actor, "email", "") or "User"


async def generate(
    db, org: OrgContext, *, actor: Any = None, committee: Any = None, meeting: Any = None,
    period_start: date | None = None, period_end: date | None = None,
    sections: Iterable[str] | None = None, title: str | None = None, reason: str = "",
):
    """Build, render and file a pack; write the ``BoardPack`` row and the audit entry.

    Refuses a bad period or section choice with :class:`PackError` before writing
    anything. Any failure after that — a query, the renderer, the disk — is kept: the
    row is written with status ``failed`` and the reason, no files are left behind, and
    the failure is in the activity trail. ``actor`` None means the scheduler did it."""
    from app.models.governance import BoardPack
    from app.services import audit, storage

    start, end = resolve_period(period_start, period_end, today=org.today, fiscal_start_month=org.fiscal_start_month)
    if sections is None and getattr(committee, "board_pack_sections", None):
        sections = committee.board_pack_sections
    keys = normalise_sections(sections, keep_order=True)
    pack_id = uuid.uuid4()
    title = (title or "").strip()[:255] or pack_title(
        committee=getattr(committee, "name", None), meeting=getattr(meeting, "title", None),
        start=start, end=end, date_format=org.date_format,
    )
    by = actor_label(actor) + (f" ({reason})" if reason else "")
    row = BoardPack(
        id=pack_id, tenant_id=org.tenant_id, committee_id=getattr(committee, "id", None),
        meeting_id=getattr(meeting, "id", None), title=title, period_start=start, period_end=end,
        sections=keys, generated_by_id=getattr(actor, "id", None), status=READY, error="",
        review_state=DRAFT, contributor_ids=[str(actor.id)] if getattr(actor, "id", None) else [],
        commentary={}, basis={}, distribution=[],
    )
    written: list[str] = []
    files: dict[str, str] = {}
    try:
        async with db.begin_nested():
            data = await build_pack(db, org, period_start=start, period_end=end, sections=keys, title=title,
                                    committee=committee, meeting=meeting, generated_by=by, viewer=actor)
            data["review_state"] = DRAFT
            row.content = encode_content(data)
            row.basis = encode_content(data.get("basis") or {})
            pdf, xlsx = to_pdf(data), to_xlsx(data)
            entity_type, entity_id = attach_target(pack_id, committee, meeting)
            stem = file_stem(title, end)
            uploader = getattr(actor, "email", "") or audit.SYSTEM_ACTOR_EMAIL
            pdf_file = await store_file(db, org.tenant_id, data=pdf, filename=f"{stem}.pdf",
                                        content_type="application/pdf", entity_type=entity_type,
                                        entity_id=entity_id, title=f"{title} (PDF)", uploaded_by=uploader)
            written.append(pdf_file.storage_key)
            xlsx_file = await store_file(db, org.tenant_id, data=xlsx, filename=f"{stem}.xlsx",
                                         content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                         entity_type=entity_type, entity_id=entity_id,
                                         title=f"{title} (XLSX)", uploaded_by=uploader)
            written.append(xlsx_file.storage_key)
            row.pdf_file_id, row.xlsx_file_id = pdf_file.id, xlsx_file.id
            files = {"pdf": pdf_file.filename, "xlsx": xlsx_file.filename}
    except Exception as exc:  # noqa: BLE001 - kept on the row, logged, never half-written
        for key in written:
            storage.delete_object(key)
        logger.exception("Board pack generation failed for tenant %s", org.tenant_id)
        row.status, row.error = FAILED, error_text(exc)
        row.pdf_file_id = row.xlsx_file_id = None
        row.content = None
        files = {}
    row.generated_at = datetime.now(timezone.utc)
    db.add(row)
    await db.flush()

    period = period_text(start, end, org.date_format)
    summary = (f"Generated board pack '{title}' for {period}" if row.status == READY
               else f"Board pack '{title}' for {period} could not be generated: {row.error}")
    changes = {
        "board_pack_id": str(pack_id), "status": row.status, "period_start": start.isoformat(),
        "period_end": end.isoformat(), "sections": keys, "files": files,
        **({"meeting": getattr(meeting, "reference", "") or getattr(meeting, "title", "")} if meeting is not None else {}),
        **({"reason": reason} if reason else {}),
    }
    entity_type, entity_id = ("committee", committee.id) if committee is not None else ("board_pack", pack_id)
    if actor is not None:
        await audit.record(db, actor=actor, action="board_pack", entity_type=entity_type, entity_id=entity_id,
                           summary=summary[:500], changes=changes)
    else:
        await audit.record_system(db, tenant_id=org.tenant_id, action="board_pack", entity_type=entity_type,
                                  entity_id=entity_id, summary=summary[:500], changes=changes)
    return row


def _lock_key(meeting_id) -> int:
    """A signed 64-bit advisory-lock key for one meeting's scheduled pack."""
    return int.from_bytes(uuid.UUID(str(meeting_id)).bytes[:8], "big", signed=True)


async def generate_due_packs(db, tenant_id, *, today: date | None = None) -> list:
    """The scheduler step: for every active committee with ``board_pack_days_before``
    set, generate the pack for each scheduled meeting that is now within that many days,
    once — a meeting that already has a pack generated inside its window is skipped (a
    pack from before the window opened is stale and does not count). A transaction-level
    advisory lock per meeting keeps two workers from generating the same pack."""
    from types import SimpleNamespace

    from sqlalchemy import func, select, text

    from app.models.governance import BoardPack, Committee, CommitteeStatus, Meeting, MeetingStatus

    if not await module_enabled(db, "governance_meetings"):
        return []
    org = await org_context(db, tenant_id)
    today = today or org.today
    committees = (await db.execute(
        select(Committee.id, Committee.name, Committee.reference, Committee.board_pack_days_before,
               Committee.board_pack_sections)
        .where(Committee.deleted.is_(False), Committee.status == CommitteeStatus.active,
               Committee.board_pack_days_before.is_not(None), Committee.board_pack_days_before > 0)
    )).all()
    made = []
    for c in committees:
        days = int(c.board_pack_days_before)
        meetings = (await db.execute(
            select(Meeting.id, Meeting.title, Meeting.reference, Meeting.meeting_date)
            .where(Meeting.committee_id == c.id, Meeting.status == MeetingStatus.scheduled,
                   Meeting.meeting_date >= today, Meeting.meeting_date <= today + timedelta(days=days))
            .order_by(Meeting.meeting_date)
        )).all()
        for m in meetings:
            if not meeting_due(m.meeting_date, days, today):
                continue
            window_at = datetime.combine(m.meeting_date - timedelta(days=days), time(0, 0), tzinfo=org.tz)
            stamped = func.coalesce(BoardPack.generated_at, BoardPack.created_at)
            if await db.scalar(select(BoardPack.id).where(BoardPack.meeting_id == m.id, stamped >= window_at).limit(1)):
                continue
            if not await db.scalar(text("SELECT pg_try_advisory_xact_lock(:k)"), {"k": _lock_key(m.id)}):
                continue
            made.append(await generate(
                db, org, committee=SimpleNamespace(id=c.id, name=c.name, reference=c.reference,
                                                   board_pack_sections=getattr(c, "board_pack_sections", None)),
                meeting=SimpleNamespace(id=m.id, title=m.title, reference=m.reference, meeting_date=m.meeting_date),
                reason=f"{days} day(s) before the meeting",
            ))
    return made


# ===========================================================================
# Phase 4B: branding, past-period snapshots, charts, stored content, sign-off
# ===========================================================================
async def load_branding(db) -> dict:
    """The organisation's pack branding (defaults when none is set). ``logo_path`` is the
    on-disk path of the logo file when one is set and still present."""
    from sqlalchemy import select

    from app.models.collab import StoredFile
    from app.models.governance import BoardPackBranding

    row = await db.scalar(select(BoardPackBranding))
    if row is None:
        return {"cover_title": "", "primary_colour": "", "classification": DEFAULT_CLASSIFICATION, "logo_path": None}
    logo_path = None
    if row.logo_file_id:
        sf = await db.scalar(select(StoredFile).where(StoredFile.id == row.logo_file_id))
        if sf is not None:
            try:
                from app.services import storage

                path = storage.resolve_path(sf.tenant_id, sf.storage_key)
                logo_path = str(path) if path.exists() else None
            except Exception:  # noqa: BLE001 - a missing logo never stops a pack
                logo_path = None
    return {"cover_title": row.cover_title or "", "primary_colour": row.primary_colour or "",
            "classification": row.classification or DEFAULT_CLASSIFICATION, "logo_path": logo_path}


async def snapshot_for(db, period_end: date) -> tuple[date, dict] | None:
    """The period snapshot nearest ``period_end`` and its figures ({key: {dimension:
    value}}), or None when none is near enough."""
    from app.services import snapshots

    dates = await snapshots.available_dates(
        db, since=period_end - timedelta(days=snapshots.MAX_DISTANCE_DAYS))
    as_of = snapshots.nearest_date(dates, period_end)
    if as_of is None:
        return None
    index = snapshots.index_rows(await snapshots.load_rows(db, [as_of]))
    return as_of, index.get(as_of, {})


def apply_snapshot(pack: dict, figures: Mapping[str, Mapping[str, dict]], as_of: date,
                   date_format: str = "DD/MM/YYYY") -> None:
    """Replace the pack's position figures with those of the period snapshot ``figures``
    (``snapshots`` keys). A figure the snapshot doesn't hold stays as at generation, and
    the section says so. Movement, failed tests and incidents are period figures and are
    never replaced. Pure (mutates ``pack``)."""
    from app.services import snapshots as sn

    when = format_date(as_of, date_format)
    notes: dict[str, list[str]] = pack.setdefault("position_notes", {})
    said = f"Position figures are from the period snapshot of {when}, the one nearest the period end."
    missing = "No period snapshot holds these figures, so they are as at the day the pack was generated."

    def one(key):
        return (figures.get(key) or {}).get("")

    if "summary" in pack:
        health = one(sn.HEALTH)
        total = one(sn.APPETITE_TOTAL)
        assurance = one(sn.ASSURANCE)
        comp = one(sn.COMPLIANCE_OVERALL)
        kri = one(sn.KRI_STATUS)
        if health or total:
            summary = pack["summary"]
            if health:
                summary["score"], summary["band"] = health.get("score"), health.get("band") or "no_data"
                summary["components"], summary["scored"], summary["total"], summary["weight_pct"] = [], 0, 0, None
            headlines = []
            if total:
                headlines.append(("Risks above tolerance", f"{total.get('breach', 0)} of {total.get('risks', 0)}"))
            if assurance:
                headlines.append(("Operating controls effective or partially effective",
                                  f"{int(assurance.get('effective', 0)) + int(assurance.get('partially_effective', 0))} "
                                  f"of {assurance.get('operating', 0)}"))
            if comp:
                headlines.append(("Clauses assured by a working control", fmt_pct(comp.get("assured_pct"))))
            if kri:
                headlines.append(("KRIs at red", str(kri.get("red", 0))))
            summary["headlines"] = headlines
            summary["actions"] = []
            notes["summary"] = [said + " The score's make-up and the open decisions are not kept in a snapshot."]
        else:
            notes["summary"] = [missing]
    if "appetite" in pack:
        total = one(sn.APPETITE_TOTAL)
        if total:
            cats = figures.get(sn.APPETITE_CATEGORY) or {}
            rows = [{"label": v.get("label") or "Category", "appetite": v.get("appetite"), "tolerance": v.get("tolerance"),
                     "risks": v.get("risks", 0), "within": v.get("within", 0), "elevated": v.get("elevated", 0),
                     "breach": v.get("breach", 0)} for _d, v in sorted(cats.items(), key=lambda kv: (kv[0] == "default", -int(kv[1].get("breach") or 0)))]
            if not rows:
                rows = [{"label": "All categories (organisation appetite)", "appetite": total.get("appetite"),
                         "tolerance": total.get("tolerance"), "risks": total.get("risks", 0),
                         "within": total.get("within", 0), "elevated": total.get("elevated", 0), "breach": total.get("breach", 0)}]
            pack["appetite"].update({"appetite": total.get("appetite"), "tolerance": total.get("tolerance"),
                                     "total": total.get("risks", 0), "within": total.get("within", 0),
                                     "elevated": total.get("elevated", 0), "breach": total.get("breach", 0), "rows": rows})
            notes["appetite"] = [said] + ([total["basis"]] if total.get("basis") else [])
        else:
            notes["appetite"] = [missing]
    if "top_risks" in pack:
        top = one(sn.TOP_RISKS)
        if top is not None:
            pack["top_risks"] = {"rows": [
                {**r, "trend_text": f"As at {when}", "trend": "unknown"} for r in top.get("rows", [])]}
            notes["top_risks"] = [said]
        else:
            notes["top_risks"] = [missing]
    if "assurance" in pack:
        a = one(sn.ASSURANCE)
        if a:
            for k in ("total", "effective", "partially_effective", "ineffective", "not_assessed", "not_operating",
                      "tests_overdue", "last_test_failed"):
                if k in a:
                    pack["assurance"][k] = a[k]
            notes["assurance"] = [said + " Failed tests are those conducted in the period."]
        else:
            notes["assurance"] = [missing]
    if "compliance" in pack:
        overall = one(sn.COMPLIANCE_OVERALL)
        if overall is not None:
            pack["compliance"] = {"overall_assured_pct": overall.get("assured_pct"), "rows": [
                {"name": v.get("name"), "applicable": v.get("applicable", 0), "assured": v.get("assured", 0),
                 "assured_pct": v.get("assured_pct"), "unassessed": v.get("unassessed", 0),
                 "failing": v.get("failing", 0), "unmapped": v.get("unmapped", 0), "gaps": v.get("gaps", 0),
                 "compliant_pct": v.get("compliant_pct")}
                for _d, v in sorted((figures.get(sn.COMPLIANCE_FRAMEWORK) or {}).items(), key=lambda kv: str(kv[1].get("name")))
            ]}
            notes["compliance"] = [said]
        else:
            notes["compliance"] = [missing]
    if "kris" in pack:
        status = one(sn.KRI_STATUS)
        if status:
            values = figures.get(sn.KRI_VALUE) or {}
            listed = sorted(
                ((d, v) for d, v in values.items() if v.get("status") in ("red", "amber")),
                key=lambda kv: (0 if kv[1].get("status") == "red" else 1, kv[1].get("reference") or "", kv[1].get("name") or ""))
            pack["kris"] = {**{k: status.get(k, 0) for k in ("green", "amber", "red", "no_data")}, "rows": [
                {"id": d, "reference": v.get("reference") or "", "name": v.get("name") or "", "status": v.get("status"),
                 "value": v.get("value"), "unit": v.get("unit") or "", "warning": None, "limit": None, "lower": None,
                 "upper": None, "direction": "", "owner": "", "as_of": as_of}
                for d, v in listed]}
            notes["kris"] = [said + " Thresholds are not kept in a snapshot."]
        else:
            notes["kris"] = [missing]
    if "issues" in pack:
        issues = one(sn.ISSUES_OPEN)
        if issues:
            by = issues.get("by_severity") or {}
            pack["issues"].update({
                "open": issues.get("open", 0), "overdue": issues.get("overdue", 0), "date_moved": 0,
                "by_severity": [{"severity": sev, "open": int(by.get(sev, 0) or 0), "overdue": 0, "date_moved": 0,
                                 "moves": 0} for sev in SEVERITIES],
                "overdue_rows": [],
            })
            notes["issues"] = [said + " The list of overdue issues and date moves are shown only for a pack as at today."]
        else:
            notes["issues"] = [missing]
    if "third_parties" in pack:
        notes["third_parties"] = [missing]


async def add_appetite_charts(db, pack: dict, org: OrgContext, period_end: date, snapshot) -> None:
    """The heat map (the snapshot's for a past period, else live) and the appetite trend
    over the last four quarter ends before the period end."""
    from sqlalchemy import select

    from app.models.risk import RiskSetting
    from app.services import snapshots
    from app.services.risk_settings import scale_for

    appetite = pack.get("appetite")
    if appetite is None:
        return
    heat = None
    if snapshot is not None:
        heat = (snapshot[1].get(snapshots.HEATMAP) or {}).get("")
    if heat is None:
        heat = await snapshots.live_heatmap(db)
    heat = dict(heat)
    settings = await db.scalar(select(RiskSetting))
    size = int(heat.get("size") or 5)
    if settings is not None:
        scale = scale_for(settings)
        heat["bands"] = {f"{like},{imp}": _v(scale.for_cell(like, imp)) for like in range(1, size + 1)
                         for imp in range(1, size + 1)}
    else:
        heat["bands"] = {f"{like},{imp}": fraction_band(like * imp, size * size)
                         for like in range(1, size + 1) for imp in range(1, size + 1)}
    appetite["heatmap"] = heat
    ends = [e for e in snapshots.quarter_ends(period_end + timedelta(days=1), TREND_QUARTERS, org.fiscal_start_month)]
    dates = await snapshots.available_dates(db, since=ends[0] - timedelta(days=snapshots.MAX_DISTANCE_DAYS))
    points = snapshots.trend_points(dates, ends)
    index = snapshots.index_rows(await snapshots.load_rows(db, [p.as_of for p in points]))
    trend = snapshots.appetite_series(index, points)
    if not trend or trend[-1]["date"] != period_end:
        trend.append({"date": period_end, "as_of": period_end, "risks": appetite.get("total"),
                      "within": appetite.get("within"), "elevated": appetite.get("elevated"),
                      "breach": appetite.get("breach")})
    appetite["trend"] = trend


def fraction_band(score: int, max_score: int) -> str:
    """Band for a cell when the organisation's bands aren't loaded. Pure."""
    share = score / max_score if max_score else 0
    return "critical" if share > 0.64 else "high" if share > 0.4 else "medium" if share > 0.16 else "low"


async def add_kri_trend(db, pack: dict, period_end: date) -> None:
    """Each listed KRI's recent readings up to the period end."""
    from sqlalchemy import select

    from app.models.operational_risk import KriMeasurement

    kris = pack.get("kris")
    if not kris:
        return
    ids = []
    for r in kris.get("rows", []):
        try:
            ids.append(uuid.UUID(str(r.get("id"))))
        except (TypeError, ValueError):
            continue
    trend: dict[str, list] = {}
    if ids:
        for kid, as_of, value in (await db.execute(
            select(KriMeasurement.kri_id, KriMeasurement.as_of_date, KriMeasurement.value)
            .where(KriMeasurement.kri_id.in_(ids), KriMeasurement.as_of_date <= period_end)
            .order_by(KriMeasurement.kri_id, KriMeasurement.as_of_date.desc(), KriMeasurement.created_at.desc())
        )).all():
            bucket = trend.setdefault(str(kid), [])
            if len(bucket) < KRI_TREND_READINGS and value is not None:
                bucket.append({"as_of": as_of, "value": float(value)})
    kris["trend"] = {k: list(reversed(v)) for k, v in trend.items()}


# ------------------------------------------------------------- stored content ---
def encode_content(value: Any) -> Any:
    """The pack's figures as JSON: dates and datetimes tagged so they come back as such;
    the timezone object is dropped (it is rebuilt from its name). Pure."""
    from decimal import Decimal

    if isinstance(value, dict):
        return {str(k): encode_content(v) for k, v in value.items() if k != "tz"}
    if isinstance(value, (list, tuple)):
        return [encode_content(v) for v in value]
    if isinstance(value, datetime):
        return {"$dt": _aware(value).isoformat()}
    if isinstance(value, date):
        return {"$d": value.isoformat()}
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "value") and not isinstance(value, (str, int, float, bool)):
        return value.value
    return value


def decode_content(value: Any) -> Any:
    """:func:`encode_content` reversed; the cover's ``tz`` is rebuilt. Pure."""
    def walk(v):
        if isinstance(v, dict):
            if set(v) == {"$dt"}:
                return datetime.fromisoformat(v["$dt"])
            if set(v) == {"$d"}:
                return date.fromisoformat(v["$d"])
            return {k: walk(x) for k, x in v.items()}
        if isinstance(v, list):
            return [walk(x) for x in v]
        return v

    out = walk(value or {})
    cover = out.get("cover")
    if isinstance(cover, dict):
        from app.services import incident_clock

        cover["tz"] = incident_clock.zone(cover.get("timezone"))
    return out


# ------------------------------------------------------------------ sign-off ---
def lifecycle_refusal(*, action: str, review_state: str, status: str, actor_id: Any,
                      contributor_ids: Iterable[Any], dual_control: bool) -> str | None:
    """Why ``actor`` may not take ``action`` (commentary | review | return | release) on a
    pack, or None. Pure.

    * commentary — not on a released pack, nor a failed one;
    * review — a ready draft; not by anyone who shaped the pack while four-eyes applies;
    * return — a reviewed pack back to draft;
    * release — a reviewed pack; not by anyone who shaped it while four-eyes applies."""
    if status != READY:
        return "This pack could not be generated; generate it again."
    if action == "commentary":
        return "A released pack is final; generate a new pack to change it." if review_state == RELEASED else None
    makers = {str(c) for c in contributor_ids or () if c}
    own = dual_control and actor_id is not None and str(actor_id) in makers
    if action == "review":
        if review_state != DRAFT:
            return f"Only a draft can be marked reviewed; this pack is {review_state}."
        if own:
            return ("Segregation of duties: you generated this pack or wrote its commentary, so someone else "
                    "must review it.")
        return None
    if action == "return":
        return None if review_state == REVIEWED else "Only a reviewed pack can be returned to draft."
    if action == "release":
        if review_state != REVIEWED:
            return "A pack must be reviewed before it is released." if review_state == DRAFT else "This pack is already released."
        if own:
            return ("Segregation of duties: you generated this pack or wrote its commentary, so someone else "
                    "must release it.")
        return None
    return f"Unknown action: {action}"


def clean_commentary(sections: Sequence[str], commentary: Mapping[str, Any]) -> dict[str, str]:
    """Commentary keyed by the pack's own sections, trimmed; blank entries dropped.
    Raises :class:`PackError` for a section the pack doesn't have or text too long. Pure."""
    out: dict[str, str] = {}
    for key, text in (commentary or {}).items():
        if key not in sections:
            raise PackError(f"This pack has no '{key}' section.")
        value = str(text or "").strip()
        if len(value) > COMMENTARY_MAX:
            raise PackError(f"Commentary for {SECTION_TITLES.get(key, key)} is longer than {COMMENTARY_MAX} characters.")
        if value:
            out[key] = value
    return out


def distribution_list(members: Iterable[Any]) -> list[dict]:
    """Who a released pack goes to: the committee's active member users, once each, in
    name order. Members need user_id, full_name, email, is_active. Pure."""
    seen: dict = {}
    for m in members:
        if not getattr(m, "is_active", True) or m.user_id in seen:
            continue
        seen[m.user_id] = {"user_id": str(m.user_id), "name": m.full_name or m.email or "", "email": m.email or "",
                           "emailed": False}
    return sorted(seen.values(), key=lambda d: d["name"].lower())


async def render_stored(db, org: OrgContext, pack, *, uploader: str) -> None:
    """Re-render the pack's files from its stored figures with its current commentary,
    review state and branding; the old files are removed once the new ones are filed."""
    from sqlalchemy import select

    from app.models.collab import StoredFile
    from app.services import storage

    if not pack.content:
        raise PackError("This pack predates stored figures; generate a new pack instead.")
    data = decode_content(pack.content)
    data["commentary"] = dict(pack.commentary or {})
    data["review_state"] = pack.review_state
    data["basis"] = decode_content(pack.basis or {}) if pack.basis else data.get("basis", {})
    data.setdefault("cover", {})["branding"] = await load_branding(db)
    pdf, xlsx = to_pdf(data), to_xlsx(data)
    entity_type, entity_id = attach_target(pack.id, _ref(pack.committee_id),
                                           _ref(pack.meeting_id))
    stem = file_stem(pack.title, pack.period_end or org.today)
    old_ids = [f for f in (pack.pdf_file_id, pack.xlsx_file_id) if f]
    pdf_file = await store_file(db, org.tenant_id, data=pdf, filename=f"{stem}.pdf", content_type="application/pdf",
                                entity_type=entity_type, entity_id=entity_id, title=f"{pack.title} (PDF)",
                                uploaded_by=uploader)
    xlsx_file = await store_file(db, org.tenant_id, data=xlsx, filename=f"{stem}.xlsx",
                                 content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 entity_type=entity_type, entity_id=entity_id, title=f"{pack.title} (XLSX)",
                                 uploaded_by=uploader)
    pack.pdf_file_id, pack.xlsx_file_id = pdf_file.id, xlsx_file.id
    if old_ids:
        for old in (await db.scalars(select(StoredFile).where(StoredFile.id.in_(old_ids)))).all():
            storage.delete_object(old.storage_key)
            await db.delete(old)
    await db.flush()


def _ref(entity_id):
    from types import SimpleNamespace

    return SimpleNamespace(id=entity_id) if entity_id else None


async def set_commentary(db, org: OrgContext, pack, actor, commentary: Mapping[str, Any]):
    """Replace the pack's commentary and re-render it. A reviewed pack goes back to draft
    (what was reviewed has changed); the writer becomes one of its makers."""
    from app.services import audit

    refusal = lifecycle_refusal(action="commentary", review_state=pack.review_state, status=pack.status,
                                actor_id=actor.id, contributor_ids=pack.contributor_ids, dual_control=False)
    if refusal:
        raise PackError(refusal)
    cleaned = clean_commentary(list(pack.sections or []), commentary)
    before = dict(pack.commentary or {})
    if cleaned == before:
        return pack
    back_to_draft = pack.review_state == REVIEWED
    pack.commentary = cleaned
    if str(actor.id) not in {str(c) for c in pack.contributor_ids or []}:
        pack.contributor_ids = list(pack.contributor_ids or []) + [str(actor.id)]
    if back_to_draft:
        pack.review_state, pack.reviewed_by_id, pack.reviewed_at = DRAFT, None, None
    await render_stored(db, org, pack, uploader=actor.email)
    changed = sorted(k for k in set(before) | set(cleaned) if before.get(k) != cleaned.get(k))
    await audit.record(
        db, actor=actor, action="board_pack_commentary", entity_type="board_pack", entity_id=pack.id,
        summary=(f"Updated commentary on board pack '{pack.title}': "
                 + ", ".join(SECTION_TITLES.get(k, k) for k in changed)
                 + ("; returned to draft for a fresh review" if back_to_draft else ""))[:500],
        changes={"sections": changed, "returned_to_draft": back_to_draft},
    )
    return pack


async def review_pack(db, pack, actor, note: str = ""):
    from app.services import audit, dual_control

    required, rule = await dual_control.dual_control_required(db, DUAL_CONTROL_MODULE, DUAL_CONTROL_ACTION)
    refusal = lifecycle_refusal(action="review", review_state=pack.review_state, status=pack.status,
                                actor_id=actor.id, contributor_ids=pack.contributor_ids, dual_control=required)
    if not refusal and required:
        refusal = await dual_control.checker_role_refusal(
            db, rule, module=DUAL_CONTROL_MODULE, action=DUAL_CONTROL_ACTION, checker_id=actor.id,
            maker_id=getattr(pack, "generated_by_id", None))
    if refusal:
        raise PackError(refusal)
    pack.review_state, pack.reviewed_by_id, pack.reviewed_at = REVIEWED, actor.id, datetime.now(timezone.utc)
    await db.flush()
    await audit.record(db, actor=actor, action="board_pack_review", entity_type="board_pack", entity_id=pack.id,
                       summary=f"Reviewed board pack '{pack.title}'" + (f": {note}" if note else ""),
                       changes={"review_state": f"{DRAFT} -> {REVIEWED}", **({"note": note} if note else {})})
    return pack


async def return_pack(db, pack, actor, note: str = ""):
    from app.services import audit

    refusal = lifecycle_refusal(action="return", review_state=pack.review_state, status=pack.status,
                                actor_id=actor.id, contributor_ids=pack.contributor_ids, dual_control=False)
    if refusal:
        raise PackError(refusal)
    pack.review_state, pack.reviewed_by_id, pack.reviewed_at = DRAFT, None, None
    await db.flush()
    await audit.record(db, actor=actor, action="board_pack_return", entity_type="board_pack", entity_id=pack.id,
                       summary=f"Returned board pack '{pack.title}' to draft" + (f": {note}" if note else ""),
                       changes={"review_state": f"{REVIEWED} -> {DRAFT}", **({"note": note} if note else {})})
    return pack


async def committee_recipients(db, committee_id) -> list[dict]:
    from sqlalchemy import select

    from app.models.governance import CommitteeMember
    from app.models.identity import User

    if committee_id is None:
        return []
    rows = (await db.execute(
        select(CommitteeMember.user_id, User.full_name, User.email, User.is_active)
        .join(User, User.id == CommitteeMember.user_id)
        .where(CommitteeMember.committee_id == committee_id)
    )).all()
    return distribution_list(rows)


async def release_pack(db, org: OrgContext, pack, actor):
    """Release a reviewed pack: final files (no draft marking), then each committee member
    user is notified in the app and e-mailed a link."""
    from app.models.enums import NotificationCategory
    from app.models.notification import EVENT_PREFIX, Notification
    from app.services import audit, dual_control, email

    required, rule = await dual_control.dual_control_required(db, DUAL_CONTROL_MODULE, DUAL_CONTROL_ACTION)
    refusal = lifecycle_refusal(action="release", review_state=pack.review_state, status=pack.status,
                                actor_id=actor.id, contributor_ids=pack.contributor_ids, dual_control=required)
    if not refusal and required:
        refusal = await dual_control.checker_role_refusal(
            db, rule, module=DUAL_CONTROL_MODULE, action=DUAL_CONTROL_ACTION, checker_id=actor.id,
            maker_id=getattr(pack, "generated_by_id", None))
    if refusal:
        raise PackError(refusal)
    pack.review_state, pack.released_by_id, pack.released_at = RELEASED, actor.id, datetime.now(timezone.utc)
    if pack.content:
        await render_stored(db, org, pack, uploader=actor.email)
    recipients = await committee_recipients(db, pack.committee_id)
    link = f"/board?pack={pack.id}"
    period = period_text(pack.period_start, pack.period_end, org.date_format) if pack.period_start and pack.period_end else ""
    for r in recipients:
        db.add(Notification(
            tenant_id=org.tenant_id, user_id=uuid.UUID(r["user_id"]), role_name="",
            title=f"Board pack released: {pack.title}"[:255],
            body=(f"The pack for {period} is ready to read." if period else "The pack is ready to read."),
            category=NotificationCategory.info, entity_type="board_pack", entity_id=pack.id, link=link,
            dedup_key=f"{EVENT_PREFIX}board-pack-released:{pack.id}:{r['user_id']}"[:255],
        ))
        if r["email"]:
            url = email.absolute_url(link)
            subject = f"{org.name}: board pack released — {pack.title}"[:200]
            text = (f"Dear {r['name']},\n\nThe board pack '{pack.title}'" + (f" for {period}" if period else "")
                    + f" has been released to you.\n\nRead it here: {url}\n\n{DEFAULT_CLASSIFICATION}.")
            html = (f"<p>Dear {escape(r['name'])},</p><p>The board pack <b>{escape(pack.title)}</b>"
                    + (f" for {escape(period)}" if period else "")
                    + f" has been released to you.</p><p><a href=\"{escape(url)}\">Open the board pack</a></p>"
                    f"<p style=\"color:#6b7280\">{DEFAULT_CLASSIFICATION}.</p>")
            try:
                r["emailed"] = bool(await email.send_email([r["email"]], subject, html, text))
            except Exception:  # noqa: BLE001 - a mail failure never undoes a release
                logger.exception("Board pack e-mail to %s failed", r["email"])
    pack.distribution = recipients
    await db.flush()
    await audit.record(
        db, actor=actor, action="board_pack_release", entity_type="board_pack", entity_id=pack.id,
        summary=(f"Released board pack '{pack.title}' to {len(recipients)} committee member(s)")[:500],
        changes={"review_state": f"{REVIEWED} -> {RELEASED}",
                 "recipients": [r["email"] or r["name"] for r in recipients],
                 "emailed": sum(1 for r in recipients if r["emailed"])},
    )
    return pack

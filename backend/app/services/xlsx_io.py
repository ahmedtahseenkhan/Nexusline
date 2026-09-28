"""Excel (.xlsx) rendering for import templates and exports.

The first sheet holds the data and is what the importer reads back (``load_xlsx``
defaults to the first sheet), so a downloaded template or export can be filled in or
edited and uploaded again unchanged. Around it:

* choice columns (enums, yes/no, a custom field's select options) get an in-cell
  dropdown, fed from a hidden ``Lists`` sheet so long option lists are not truncated;
* required headers are shaded and carry a "Required" note; help text becomes a note;
* a ``Guide`` sheet lists every column — required, type, allowed values, example.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timezone
from enum import Enum
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if TYPE_CHECKING:
    from app.services.import_registry import Column

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_REQUIRED_FILL = PatternFill("solid", fgColor="B45309")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_BOOL_OPTIONS = ["true", "false"]
_DATA_ROWS = 1000  # dropdowns reach this far down the data sheet
_KIND_LABEL = {
    "text": "Text",
    "int": "Whole number",
    "float": "Number",
    "bool": "Yes / no (true or false)",
    "date": "Date (YYYY-MM-DD)",
    "enum": "One of the listed values",
    "link": "Name or reference of an existing record",
}


def _sheet_title(name: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Data"


def _options(col: "Column") -> list[str]:
    if col.kind == "enum":
        return list(col.enum_values or [])
    if col.kind == "bool":
        return _BOOL_OPTIONS
    return []


def _cell_value(value: object) -> object:
    """Values openpyxl writes natively stay typed (so Excel sees real dates/numbers)."""
    if value is None:
        return None
    if isinstance(value, Enum):
        value = value.value
    if isinstance(value, datetime) and value.tzinfo is not None:
        # Excel has no time zones: write the instant in UTC.
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    if isinstance(value, (bool, int, float, date, datetime)):
        return value
    return str(value)


def build_workbook(
    columns: list["Column"],
    rows: list[dict],
    *,
    title: str,
    examples: dict[str, str] | None = None,
) -> bytes:
    """Render ``rows`` (header -> value) under ``columns`` as an .xlsx workbook."""
    wb = Workbook()
    data = wb.active
    data.title = _sheet_title(title)
    lists = wb.create_sheet("Lists")
    guide = wb.create_sheet("Guide")

    list_col = 0
    for idx, col in enumerate(columns, start=1):
        cell = data.cell(row=1, column=idx, value=col.header)
        cell.font = _HEADER_FONT
        cell.fill = _REQUIRED_FILL if col.required else _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        note = "\n".join(p for p in ("Required" if col.required else "", col.help) if p)
        if note:
            cell.comment = Comment(note, "GRC")
        letter = get_column_letter(idx)
        data.column_dimensions[letter].width = max(14, min(40, len(col.header) + 4))

        options = _options(col)
        if options:
            list_col += 1
            list_letter = get_column_letter(list_col)
            lists.cell(row=1, column=list_col, value=col.header)
            for r, option in enumerate(options, start=2):
                lists.cell(row=r, column=list_col, value=option)
            dv = DataValidation(
                type="list",
                formula1=f"=Lists!${list_letter}$2:${list_letter}${len(options) + 1}",
                allow_blank=not col.required,
                showErrorMessage=True,
                errorTitle="Not an allowed value",
                error=f"Pick a value from the list for '{col.header}'.",
            )
            dv.add(f"{letter}2:{letter}{_DATA_ROWS + 1}")
            data.add_data_validation(dv)

    for r, row in enumerate(rows, start=2):
        for idx, col in enumerate(columns, start=1):
            value = _cell_value(row.get(col.header))
            if value is None:
                continue
            cell = data.cell(row=r, column=idx, value=value)
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
    data.freeze_panes = "A2"
    lists.sheet_state = "hidden"

    guide.append(["Column", "Required", "Type", "Allowed values", "Example", "Notes"])
    for cell in guide[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for col in columns:
        guide.append([
            col.header,
            "Yes" if col.required else "",
            _KIND_LABEL.get(col.kind, col.kind),
            ", ".join(_options(col)),
            (examples or {}).get(col.header, ""),
            col.help,
        ])
    for letter, width in zip("ABCDEF", (32, 10, 30, 50, 24, 60)):
        guide.column_dimensions[letter].width = width
    guide.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
